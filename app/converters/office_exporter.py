from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import tempfile
import re
from io import BytesIO

import fitz


EMU_PER_INCH = 914400
PT_PER_INCH = 72.0

_KIWI_SPACING_ENGINE = None
_KIWI_SPACING_UNAVAILABLE = False



@dataclass
class OfficeExportOptions:
    dpi: int = 200
    page_range: list[int] | None = None
    fit_mode: str = "fit"
    include_page_breaks: bool = True

    # PowerPoint only:
    # If True, PDF text spans are added as editable PowerPoint text boxes
    # on top of the page image.
    include_editable_text_layer: bool = False
    # PowerPoint mode:
    # - auto: choose image/text mode per page
    # - image: page image only
    # - text_only: editable text + embedded image/drawing objects, no full-page image
    powerpoint_mode: str = "image"

    # Drawing conversion safety limit.
    # Complex PDFs can contain thousands of vector paths and freeze the UI.
    max_drawing_objects: int = 1200

    # Faded background mode options.
    # Higher background transparency = clearer original PDF background.
    faded_background_transparency: float = 78.0

    # Lower text box transparency = stronger white mask behind editable text.
    text_box_fill_transparency: float = 18.0


def parse_page_range(range_text: str, page_count: int) -> list[int]:
    """
    Convert user page range like "1-3,5,8-10" to zero-based page indexes.
    Empty text means all pages.
    """
    text = (range_text or "").strip()
    if not text:
        return list(range(page_count))

    result: list[int] = []
    seen: set[int] = set()

    for part in text.split(","):
        part = part.strip()
        if not part:
            continue

        if "-" in part:
            left, right = part.split("-", 1)
            start = int(left.strip())
            end = int(right.strip())

            if start > end:
                start, end = end, start

            for page_no in range(start, end + 1):
                index = page_no - 1
                if 0 <= index < page_count and index not in seen:
                    result.append(index)
                    seen.add(index)
        else:
            page_no = int(part)
            index = page_no - 1
            if 0 <= index < page_count and index not in seen:
                result.append(index)
                seen.add(index)

    return result



def _pdf_rgb_float_to_tuple(color_value) -> tuple[int, int, int]:
    """
    Convert PyMuPDF color tuple like (0.1, 0.2, 0.3) to RGB 0-255.
    """
    try:
        if color_value is None:
            return 0, 0, 0

        if isinstance(color_value, (list, tuple)) and len(color_value) >= 3:
            r = int(max(0, min(1, float(color_value[0]))) * 255)
            g = int(max(0, min(1, float(color_value[1]))) * 255)
            b = int(max(0, min(1, float(color_value[2]))) * 255)
            return r, g, b
    except Exception:
        pass

    return 0, 0, 0


def _pdf_color_int_to_rgb(color_value) -> tuple[int, int, int]:
    """
    PyMuPDF span color is usually an integer like 0xRRGGBB.
    """
    try:
        value = int(color_value)
        r = (value >> 16) & 255
        g = (value >> 8) & 255
        b = value & 255
        return r, g, b
    except Exception:
        return 0, 0, 0


def _safe_text(text: str) -> str:
    return _fix_broken_pdf_symbols((text or "").replace("\x00", "")).strip()


def _contains_cjk(text: str) -> bool:
    """Return True if text contains Korean/Japanese/Chinese characters."""
    for ch in text or "":
        code = ord(ch)
        if (
            0xAC00 <= code <= 0xD7AF  # Hangul syllables
            or 0x1100 <= code <= 0x11FF  # Hangul Jamo
            or 0x3130 <= code <= 0x318F  # Hangul compatibility Jamo
            or 0x3040 <= code <= 0x30FF  # Japanese Hiragana/Katakana
            or 0x4E00 <= code <= 0x9FFF  # CJK Unified Ideographs
        ):
            return True
    return False


def _normalize_line_text(text: str) -> str:
    """Clean spacing for PowerPoint text boxes."""
    text = (text or "").replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _join_spans_with_estimated_spaces(spans: list[dict]) -> str:
    """
    Join PDF spans into one visual line.

    PDF text often arrives as many small spans. If we simply concatenate them,
    words can stick together. This function estimates spaces from horizontal gaps.
    """
    parts: list[str] = []
    previous_x1: float | None = None
    previous_size: float = 10.0

    for span in spans:
        text = _safe_text(span.get("text", ""))
        if not text:
            continue

        bbox = span.get("bbox")
        if not bbox or len(bbox) != 4:
            continue

        x0, _y0, x1, _y1 = [float(v) for v in bbox]
        size = float(span.get("size", previous_size) or previous_size)

        if previous_x1 is not None:
            gap = x0 - previous_x1
            # Approximate one space width. CJK documents often have smaller gaps.
            one_space = max(size * 0.25, 2.0)

            if gap > one_space:
                space_count = max(1, min(6, round(gap / max(size * 0.45, 3.0))))
                if parts and not parts[-1].endswith(" "):
                    parts.append(" " * space_count)

        parts.append(text)
        previous_x1 = x1
        previous_size = size

    return _apply_korean_spacing_heuristics(_normalize_line_text("".join(parts)))




def _apply_korean_spacing_heuristics(text: str) -> str:
    """
    Best-effort Korean spacing restoration for PDF text extraction.

    Important:
    - This is not a full Korean NLP tokenizer.
    - It fixes common no-space extraction patterns often found in manuals.
    - It is intentionally conservative to avoid breaking technical terms too much.
    """
    text = text or ""

    # CJK/ASCII boundary spacing.
    text = re.sub(
        r"([\uAC00-\uD7AF\u3130-\u318F\u1100-\u11FF\u3040-\u30FF\u4E00-\u9FFF])([A-Za-z0-9])",
        r"\1 \2",
        text,
    )
    text = re.sub(
        r"([A-Za-z0-9])([\uAC00-\uD7AF\u3130-\u318F\u1100-\u11FF\u3040-\u30FF\u4E00-\u9FFF])",
        r"\1 \2",
        text,
    )

    # Common Korean manual/product-document phrases.
    replacements = [
        ("지구온난화방지를위하여", "지구 온난화 방지를 위하여"),
        ("지구온난화방지", "지구 온난화 방지"),
        ("방지를위하여", "방지를 위하여"),
        ("본제품은", "본 제품은"),
        ("본제품", "본 제품"),
        ("재활용및재사용이가능한", "재활용 및 재사용이 가능한"),
        ("재활용및재사용", "재활용 및 재사용"),
        ("재사용이가능한", "재사용이 가능한"),
        ("가능한포장재를", "가능한 포장재를"),
        ("포장재를사용합니다", "포장재를 사용합니다"),
        ("자세한제품정보는", "자세한 제품 정보는"),
        ("자세한제품정보", "자세한 제품 정보"),
        ("제품정보는", "제품 정보는"),
        ("제품정보", "제품 정보"),
        ("웹사이트를참조하십시오", "웹사이트를 참조하십시오"),
        ("웹사이트를참조", "웹사이트를 참조"),
        ("참조하십시오", "참조하십시오"),
    ]

    for old, new in replacements:
        text = text.replace(old, new)

    # Generic small rules for common particles/phrases in manuals.
    generic_patterns = [
        (r"([가-힣]{2,})(및)([가-힣]{2,})", r"\1 및 \3"),
        (r"([가-힣]{2,})(또는)([가-힣]{2,})", r"\1 또는 \3"),
        (r"([가-힣]{2,})(그리고)([가-힣]{2,})", r"\1 그리고 \3"),
        (r"([가-힣]{2,})(위하여)", r"\1 위하여"),
        (r"([가-힣]{2,})(가능한)", r"\1 가능한"),
        (r"([가-힣]{2,})(참조하십시오)", r"\1 참조하십시오"),
        (r"([가-힣]{2,})(사용합니다)", r"\1 사용합니다"),
        (r"([가-힣]{2,})(확인하십시오)", r"\1 확인하십시오"),
    ]

    for pattern, repl in generic_patterns:
        text = re.sub(pattern, repl, text)

    # Cleanup spacing around punctuation and repeated spaces.
    text = re.sub(r"\s+([.,;:!?%)]}）】、。！？；：])", r"\1", text)
    text = re.sub(r"([([{（【])\s+", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text)

    return _final_korean_spacing_cleanup(text).strip()



def _final_korean_spacing_cleanup(text: str) -> str:
    """
    Final forced cleanup for partially spaced Korean text.
    This runs after span joining, rule spacing, and Kiwi spacing.
    """
    text = text or ""

    replacements = [
        # Exact problematic patterns from user PDF.
        ("지구 온난화 방지를 위하여본 제품은재활용 및 재사용이 가능한 포장재를",
         "지구 온난화 방지를 위하여 본 제품은 재활용 및 재사용이 가능한 포장재를"),
        ("GIGABYTE 는소비자와함께환경을보호합니다",
         "GIGABYTE는 소비자와 함께 환경을 보호합니다"),
        ("GIGABYTE는소비자와함께환경을보호합니다",
         "GIGABYTE는 소비자와 함께 환경을 보호합니다"),

        # Partial glued fragments.
        ("위하여본", "위하여 본"),
        ("제품은재활용", "제품은 재활용"),
        ("제품은 재활용", "제품은 재활용"),
        ("는소비자와", "는 소비자와"),
        ("은소비자와", "은 소비자와"),
        ("소비자와함께", "소비자와 함께"),
        ("함께환경을", "함께 환경을"),
        ("환경을보호합니다", "환경을 보호합니다"),
        ("재사용이가능한", "재사용이 가능한"),
        ("가능한포장재를", "가능한 포장재를"),
        ("포장재를사용하였습니다", "포장재를 사용하였습니다"),

        # Common manual phrases.
        ("자세한제품정보는", "자세한 제품 정보는"),
        ("제품정보는", "제품 정보는"),
        ("웹사이트를참조하십시오", "웹사이트를 참조하십시오"),
        ("웹사이트를 참조하십시오", "웹사이트를 참조하십시오"),
    ]

    for old, new in replacements:
        text = text.replace(old, new)

    # Brand + Korean particle cleanup.
    text = re.sub(r"\b(GIGABYTE)\s+(은|는|이|가|을|를|와|과|의)", r"\1\2", text)

    # Add missing spaces after common Korean endings when followed by Korean noun phrases.
    text = re.sub(r"(위하여)(본)", r"\1 \2", text)
    text = re.sub(r"(제품은)(재활용)", r"\1 \2", text)
    text = re.sub(r"(소비자와)(함께)", r"\1 \2", text)
    text = re.sub(r"(함께)(환경을)", r"\1 \2", text)
    text = re.sub(r"(환경을)(보호합니다)", r"\1 \2", text)

    # Cleanup spaces.
    text = re.sub(r"\s+([.,;:!?%)]}）】、。！？；：])", r"\1", text)
    text = re.sub(r"([([{（【])\s+", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text)

    return text.strip()




def _has_long_korean_run(text: str) -> bool:
    return bool(re.search(r"[가-힣]{6,}", text or ""))


def _space_known_manual_phrases(text: str) -> str:
    """
    Common Korean manual/legal/product notice phrase spacing.
    This runs before and after Kiwi because PDF extraction can glue words badly.
    """
    text = text or ""

    replacements = [
        # New user-provided examples
        ("이설명서에포함된정보는저작권법보호를받으며", "이 설명서에 포함된 정보는 저작권법 보호를 받으며"),
        ("이설명서에포함된정보는", "이 설명서에 포함된 정보는"),
        ("이설명서", "이 설명서"),
        ("설명서에포함된", "설명서에 포함된"),
        ("포함된정보는", "포함된 정보는"),
        ("저작권법보호를받으며", "저작권법 보호를 받으며"),
        ("저작권법보호", "저작권법 보호"),
        ("보호를받으며", "보호를 받으며"),
        ("GIGABYTE의재산입니다", "GIGABYTE의 재산입니다"),
        ("의재산입니다", "의 재산입니다"),
        ("사전서면허가없이는", "사전 서면 허가 없이는"),
        ("사전서면허가", "사전 서면 허가"),
        ("허가없이는", "허가 없이는"),
        ("이문서의일부", "이 문서의 일부"),
        ("이문서", "이 문서"),
        ("문서의일부", "문서의 일부"),
        ("또는전부를", "또는 전부를"),
        ("전부를어떤", "전부를 어떤"),

        # Previous examples
        ("지구온난화방지를위하여본제품은재활용및재사용이가능한포장재를사용하였습니다", "지구 온난화 방지를 위하여 본 제품은 재활용 및 재사용이 가능한 포장재를 사용하였습니다"),
        ("지구온난화방지를위하여", "지구 온난화 방지를 위하여"),
        ("지구온난화방지", "지구 온난화 방지"),
        ("위하여본", "위하여 본"),
        ("본제품은", "본 제품은"),
        ("제품은재활용", "제품은 재활용"),
        ("재활용및재사용", "재활용 및 재사용"),
        ("재사용이가능한", "재사용이 가능한"),
        ("가능한포장재를", "가능한 포장재를"),
        ("포장재를사용하였습니다", "포장재를 사용하였습니다"),
        ("소비자와함께환경을보호합니다", "소비자와 함께 환경을 보호합니다"),
        ("는소비자와", "는 소비자와"),
        ("소비자와함께", "소비자와 함께"),
        ("함께환경을", "함께 환경을"),
        ("환경을보호합니다", "환경을 보호합니다"),
        ("자세한제품정보는", "자세한 제품 정보는"),
        ("제품정보는", "제품 정보는"),
        ("웹사이트를참조하십시오", "웹사이트를 참조하십시오"),
    ]

    for old, new in replacements:
        text = text.replace(old, new)

    return text


def _protect_brand_tokens(text: str) -> tuple[str, dict[str, str]]:
    """
    Protect brand/ASCII tokens before aggressive Korean spacing.
    """
    tokens: dict[str, str] = {}

    def repl(match):
        key = f"__BRANDTOKEN{len(tokens)}__"
        tokens[key] = match.group(0)
        return f" {key} "

    protected = re.sub(r"[A-Za-z][A-Za-z0-9._+\-/]*", repl, text or "")
    return protected, tokens


def _restore_brand_tokens(text: str, tokens: dict[str, str]) -> str:
    for key, value in tokens.items():
        text = text.replace(key, value)

    # Brand + particle cleanup
    text = re.sub(r"\b([A-Za-z][A-Za-z0-9._+\-/]*)\s+(은|는|이|가|을|를|와|과|의)", r"\1\2", text)
    text = re.sub(r"\b([A-Za-z][A-Za-z0-9._+\-/]*)(의)([가-힣])", r"\1\2 \3", text)
    return text



def _get_kiwi_spacing_engine():
    """
    Lazy-load kiwipiepy Kiwi spacing engine.
    If unavailable, return None and keep app working.
    """
    global _KIWI_SPACING_ENGINE, _KIWI_SPACING_UNAVAILABLE

    try:
        if _KIWI_SPACING_UNAVAILABLE:
            return None
    except NameError:
        _KIWI_SPACING_UNAVAILABLE = False

    try:
        if _KIWI_SPACING_ENGINE is not None:
            return _KIWI_SPACING_ENGINE
    except NameError:
        _KIWI_SPACING_ENGINE = None

    try:
        from kiwipiepy import Kiwi
        _KIWI_SPACING_ENGINE = Kiwi()
        return _KIWI_SPACING_ENGINE
    except Exception:
        _KIWI_SPACING_UNAVAILABLE = True
        return None


def _aggressive_korean_spacing(text: str) -> str:
    """
    Aggressive Korean spacing for long PDF-extracted Korean lines.

    Strategy:
    - Apply known phrase spacing first
    - Protect ASCII/brand tokens
    - For long Korean runs, remove unreliable spaces inside Korean chunks
    - Run Kiwi.space if available
    - Restore brand tokens
    - Apply final cleanup
    """
    text = text or ""
    text = re.sub(r"\s{2,}", " ", text).strip()

    if not text or not _contains_cjk(text):
        return text

    text = _space_known_manual_phrases(text)

    protected, tokens = _protect_brand_tokens(text)

    # If there are long glued Korean chunks, make existing Korean spacing less trusted.
    # Keep spaces around protected brand tokens.
    if _has_long_korean_run(protected):
        # Remove spaces only when both sides are Korean to let Kiwi re-space the chunk.
        protected = re.sub(r"([가-힣])\s+([가-힣])", r"\1\2", protected)

    try:
        kiwi = _get_kiwi_spacing_engine()
    except Exception:
        kiwi = None

    if kiwi is not None:
        try:
            protected = kiwi.space(protected)
        except Exception:
            pass

    text = _restore_brand_tokens(protected, tokens)
    text = _space_known_manual_phrases(text)

    # Sentence/brand boundary cleanup.
    text = re.sub(r"(입니다)([A-Za-z])", r"\1 \2", text)
    text = re.sub(r"(합니다)([A-Za-z])", r"\1 \2", text)
    text = re.sub(r"(십시오)([A-Za-z])", r"\1 \2", text)

    # Common Korean glued endings.
    cleanup_patterns = [
        (r"(정보는)(저작권법)", r"\1 \2"),
        (r"(받으며)([A-Za-z])", r"\1 \2"),
        (r"(입니다)([A-Za-z])", r"\1 \2"),
        (r"(없이는)(이)", r"\1 \2"),
        (r"(일부)(또는)", r"\1 또는"),
        (r"(전부를)(어떤)", r"\1 \2"),
        (r"(재산입니다)(GIGABYTE)", r"\1 \2"),
    ]

    for pattern, repl in cleanup_patterns:
        text = re.sub(pattern, repl, text)

    # Punctuation and spacing cleanup.
    text = re.sub(r"\s+([.,;:!?%)]}）】、。！？；：])", r"\1", text)
    text = re.sub(r"([([{（【])\s+", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text)

    return text.strip()




# --- forced brand token override v3 ---
def _protect_brand_tokens(text: str) -> tuple[str, dict[str, str]]:
    """
    Protect ASCII/brand tokens before Korean spacing.

    Uses compact markers that are less likely to be broken by Kiwi.
    """
    tokens: dict[str, str] = {}

    def repl(match):
        key = f"§BT{len(tokens)}§"
        tokens[key] = match.group(0)
        return f" {key} "

    protected = re.sub(r"[A-Za-z][A-Za-z0-9._+\-/]*", repl, text or "")
    return protected, tokens


def _restore_brand_tokens(text: str, tokens: dict[str, str]) -> str:
    """
    Restore protected tokens robustly.

    Supports:
    - §BT0§
    - § BT0 §
    - §BT 0§
    - leaked old __BRANDTOKEN0__ variants
    """
    text = text or ""

    for key, value in tokens.items():
        text = text.replace(key, value)

        m = re.search(r"BT(\d+)", key)
        if m:
            idx = m.group(1)

            variants = [
                rf"§\s*BT\s*{idx}\s*§",
                rf"§BT\s*{idx}§",
                rf"§\s*BT{idx}\s*§",
                rf"§BT{idx}§",
            ]

            for pattern in variants:
                text = re.sub(pattern, value, text)

    # Remove leaked old placeholders.
    text = re.sub(r"__\s*BRANDTOKEN\s*\d+\s*__", "", text)
    text = re.sub(r"__BRANDTOKEN\d+\s*__", "", text)
    text = re.sub(r"__\s*BRANDTOKEN\d+\s*__", "", text)
    text = re.sub(r"__BRANDTOKEN\d+\s*_", "", text)
    text = re.sub(r"_\s*BRANDTOKEN\d+\s*__", "", text)

    # Remove leaked new placeholders.
    text = re.sub(r"§\s*BT\s*\d+\s*§", "", text)

    # Brand + Korean particle cleanup.
    text = re.sub(
        r"\b([A-Za-z][A-Za-z0-9._+\-/]*)\s+(은|는|이|가|을|를|와|과|의)",
        r"\1\2",
        text,
    )
    text = re.sub(
        r"\b([A-Za-z][A-Za-z0-9._+\-/]*)(의)([가-힣])",
        r"\1\2 \3",
        text,
    )

    # Quote/punctuation cleanup.
    text = re.sub(r'"\s+', '"', text)
    text = re.sub(r'\s+"', '"', text)
    text = re.sub(r"\s+([:;,.!?])", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text)

    return text.strip()


def _remove_leaked_brand_tokens(text: str) -> str:
    """
    Final guard against leaked placeholder tokens.
    """
    text = text or ""

    # Old placeholder variants.
    text = re.sub(r"__\s*BRANDTOKEN\s*\d+\s*__", "", text)
    text = re.sub(r"__BRANDTOKEN\d+\s*__", "", text)
    text = re.sub(r"__\s*BRANDTOKEN\d+\s*__", "", text)
    text = re.sub(r"__BRANDTOKEN\d+\s*_", "", text)
    text = re.sub(r"_\s*BRANDTOKEN\d+\s*__", "", text)
    text = re.sub(r"BRANDTOKEN\s*\d+", "", text)

    # New placeholder variants.
    text = re.sub(r"§\s*BT\s*\d+\s*§", "", text)

    # Clean awkward spaces.
    text = re.sub(r'"\s+', '"', text)
    text = re.sub(r'\s+"', '"', text)
    text = re.sub(r"\s+([:;,.!?])", r"\1", text)
    text = re.sub(r"\s{2,}", " ", text)

    return text.strip()




def _fix_broken_pdf_symbols(text: str) -> str:
    """
    Fix common broken symbol-font characters extracted from PDFs.

    Some PDFs store checkbox/square symbols using custom fonts.
    PyMuPDF may extract them as control/private characters such as ,
    which PowerPoint can display as an incorrect number-like glyph.
    """
    text = text or ""

    replacements = {
        "\x84": "□",
        "": "□",
        "\uf0a3": "□",
        "\uf0fe": "□",
        "\uf06f": "○",
        "\uf0b7": "•",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    # Remove remaining invisible C0/C1 control characters except common whitespace.
    text = "".join(
        ch for ch in text
        if ch in ("\t", "\n", "\r") or ord(ch) >= 32 and not (0x7F <= ord(ch) <= 0x9F)
    )

    return text



class OfficeExporter:
    """
    Local PDF -> Office exporter.

    PowerPoint:
      - high quality image slide mode
      - optional editable text layer
      - preserves visual appearance well

    Word:
      - extracts plain text page by page

    Excel:
      - extracts text lines into rows
    """

    def export_pdf_to_powerpoint(
        self,
        source_pdf_path: str,
        output_pptx_path: str,
        options: OfficeExportOptions | None = None,
        progress_callback=None,
    ) -> None:
        try:
            from pptx import Presentation
            from pptx.util import Emu
        except ImportError as exc:
            raise RuntimeError(
                "python-pptx is not installed. Run: pip install python-pptx Pillow"
            ) from exc

        options = options or OfficeExportOptions()
        source_path = Path(source_pdf_path)
        output_path = Path(output_pptx_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = fitz.open(str(source_path))

        try:
            page_indexes = options.page_range or list(range(len(doc)))

            if not page_indexes:
                raise RuntimeError("No pages selected for PowerPoint export.")

            first_page = doc.load_page(page_indexes[0])
            first_rect = first_page.rect

            prs = Presentation()
            prs.slide_width = Emu(int((first_rect.width / PT_PER_INCH) * EMU_PER_INCH))
            prs.slide_height = Emu(int((first_rect.height / PT_PER_INCH) * EMU_PER_INCH))

            blank_layout = prs.slide_layouts[6]
            zoom = max(options.dpi, 72) / 72.0

            with tempfile.TemporaryDirectory(prefix="pdf_pptx_") as temp_dir:
                temp_dir_path = Path(temp_dir)

                total_pages = len(page_indexes)

                for output_no, page_index in enumerate(page_indexes, start=1):
                    if progress_callback:
                        should_continue = progress_callback(
                            output_no - 1,
                            total_pages,
                            f"페이지 {output_no}/{total_pages} 준비 중..."
                        )
                        if should_continue is False:
                            raise RuntimeError("PowerPoint conversion was cancelled.")

                    if page_index < 0 or page_index >= len(doc):
                        continue

                    page = doc.load_page(page_index)
                    page_rect = page.rect

                    if progress_callback:
                        should_continue = progress_callback(
                            output_no - 1,
                            total_pages,
                            f"페이지 {output_no}/{total_pages} 렌더링 중... 큰 페이지는 시간이 걸릴 수 있습니다."
                        )
                        if should_continue is False:
                            raise RuntimeError("PowerPoint conversion was cancelled.")

                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(zoom, zoom),
                        alpha=False,
                        colorspace=fitz.csRGB,
                    )

                    image_path = temp_dir_path / f"page_{output_no:04d}.png"
                    pix.save(str(image_path))

                    slide = prs.slides.add_slide(blank_layout)

                    slide_w = prs.slide_width
                    slide_h = prs.slide_height

                    page_ratio = page_rect.width / page_rect.height
                    slide_ratio = int(slide_w) / int(slide_h)

                    if options.fit_mode == "fill":
                        if page_ratio > slide_ratio:
                            image_height = slide_h
                            image_width = int(image_height * page_ratio)
                        else:
                            image_width = slide_w
                            image_height = int(image_width / page_ratio)
                    else:
                        if page_ratio > slide_ratio:
                            image_width = slide_w
                            image_height = int(image_width / page_ratio)
                        else:
                            image_height = slide_h
                            image_width = int(image_height * page_ratio)

                    image_left = int((int(slide_w) - int(image_width)) / 2)
                    image_top = int((int(slide_h) - int(image_height)) / 2)

                    powerpoint_mode = getattr(options, "powerpoint_mode", "image")

                    if powerpoint_mode != "text_only":
                        slide.shapes.add_picture(
                            str(image_path),
                            image_left,
                            image_top,
                            width=image_width,
                            height=image_height,
                        )

                        if powerpoint_mode == "faded_image_text":
                            self._add_faded_background_overlay(
                                slide=slide,
                                image_left=image_left,
                                image_top=image_top,
                                image_width=int(image_width),
                                image_height=int(image_height),
                                transparency=45.0,
                            )

                    if powerpoint_mode == "text_only":
                        if progress_callback:
                            should_continue = progress_callback(
                                output_no - 1,
                                total_pages,
                                f"페이지 {output_no}/{total_pages} 도형/이미지 변환 중..."
                            )
                            if should_continue is False:
                                raise RuntimeError("PowerPoint conversion was cancelled.")

                        self._add_pdf_drawings_to_slide(
                            slide=slide,
                            page=page,
                            image_left=image_left,
                            image_top=image_top,
                            image_width=int(image_width),
                            image_height=int(image_height),
                            max_drawing_objects=getattr(options, "max_drawing_objects", 800),
                        )

                        self._add_pdf_image_blocks_to_slide(
                            slide=slide,
                            page=page,
                            image_left=image_left,
                            image_top=image_top,
                            image_width=int(image_width),
                            image_height=int(image_height),
                        )

                    if options.include_editable_text_layer or powerpoint_mode in ("text_only", "faded_image_text"):
                        self._current_text_box_fill_transparency = getattr(
                            options,
                            "text_box_fill_transparency",
                            18.0,
                        )

                        if progress_callback:
                            should_continue = progress_callback(
                                output_no - 1,
                                total_pages,
                                f"페이지 {output_no}/{total_pages} 텍스트 변환 중..."
                            )
                            if should_continue is False:
                                raise RuntimeError("PowerPoint conversion was cancelled.")

                        self._add_editable_text_layer_to_slide(
                            slide=slide,
                            page=page,
                            image_left=image_left,
                            image_top=image_top,
                            image_width=int(image_width),
                            image_height=int(image_height),
                        )

                    if progress_callback:
                        should_continue = progress_callback(
                            output_no,
                            total_pages,
                            f"페이지 {output_no}/{total_pages} 완료"
                        )
                        if should_continue is False:
                            raise RuntimeError("PowerPoint conversion was cancelled.")

            prs.save(str(output_path))

        finally:
            doc.close()

    def _add_faded_background_overlay(
        self,
        slide,
        image_left: int,
        image_top: int,
        image_width: int,
        image_height: int,
        transparency: float = 45.0,
    ) -> None:
        """
        Add a semi-transparent white overlay on top of the page image.

        python-pptx's high-level transparency API is not reliable across versions,
        so this writes the PowerPoint XML alpha value directly.

        transparency:
          0   = solid white
          100 = fully transparent
        """
        try:
            from pptx.util import Emu
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.dml.color import RGBColor
            from pptx.oxml.xmlchemy import OxmlElement
            from pptx.oxml.ns import qn
        except ImportError:
            return

        try:
            overlay = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Emu(image_left),
                Emu(image_top),
                Emu(image_width),
                Emu(image_height),
            )

            overlay.fill.solid()
            overlay.fill.fore_color.rgb = RGBColor(255, 255, 255)
            overlay.line.fill.background()

            # DrawingML alpha is opacity, not transparency.
            # 100000 = fully opaque, 0 = fully transparent.
            transparency = max(0.0, min(100.0, float(transparency)))
            opacity = int((100.0 - transparency) * 1000)

            sp_pr = overlay._element.spPr
            solid_fill = sp_pr.find(qn("a:solidFill"))

            if solid_fill is not None:
                srgb = solid_fill.find(qn("a:srgbClr"))

                if srgb is not None:
                    # Remove old alpha nodes if any.
                    for old_alpha in list(srgb.findall(qn("a:alpha"))):
                        srgb.remove(old_alpha)

                    alpha = OxmlElement("a:alpha")
                    alpha.set("val", str(opacity))
                    srgb.append(alpha)

            # Put overlay behind editable text but above page image.
            return

        except Exception:
            return


    def _add_pdf_drawings_to_slide(
        self,
        slide,
        page: fitz.Page,
        image_left: int,
        image_top: int,
        image_width: int,
        image_height: int,
        max_drawing_objects: int = 1200,
    ) -> None:
        """
        Add simple PDF vector drawings as PowerPoint shapes.

        Supported:
        - straight lines
        - rectangles

        Not yet supported:
        - curves
        - complex paths
        - clipping masks
        - gradients
        """
        try:
            from pptx.util import Emu, Pt
            from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
            from pptx.dml.color import RGBColor
        except ImportError as exc:
            raise RuntimeError("python-pptx is required for drawing export.") from exc

        page_rect = page.rect
        page_width = max(float(page_rect.width), 1.0)
        page_height = max(float(page_rect.height), 1.0)

        scale_x = image_width / page_width
        scale_y = image_height / page_height

        def map_x(x: float) -> int:
            return image_left + int(float(x) * scale_x)

        def map_y(y: float) -> int:
            return image_top + int(float(y) * scale_y)

        def map_w(w: float) -> int:
            return max(int(float(w) * scale_x), 1)

        def map_h(h: float) -> int:
            return max(int(float(h) * scale_y), 1)

        try:
            drawings = page.get_drawings()
        except Exception:
            return

        # Safety guard:
        # Some PDFs contain thousands of tiny vector paths.
        # Converting all of them can make PowerPoint generation extremely slow.
        if len(drawings) > max_drawing_objects:
            return

        for drawing in drawings:
            try:
                stroke_color = drawing.get("color")
                fill_color = drawing.get("fill")
                line_width = float(drawing.get("width", 0.5) or 0.5)

                stroke_rgb = _pdf_rgb_float_to_tuple(stroke_color)
                fill_rgb = _pdf_rgb_float_to_tuple(fill_color)

                # 1) Rectangle-like drawing.
                rect = drawing.get("rect")
                items = drawing.get("items", [])

                is_rect = False
                if rect is not None:
                    # PyMuPDF rect can represent bounding box of a path.
                    # Treat as rectangle only if path has rectangle item or fill exists.
                    for item in items:
                        if item and item[0] == "re":
                            is_rect = True
                            break

                    if fill_color is not None:
                        is_rect = True

                if is_rect and rect is not None:
                    x0, y0, x1, y1 = float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)

                    shape = slide.shapes.add_shape(
                        MSO_SHAPE.RECTANGLE,
                        Emu(map_x(x0)),
                        Emu(map_y(y0)),
                        Emu(map_w(x1 - x0)),
                        Emu(map_h(y1 - y0)),
                    )

                    # Fill
                    if fill_color is not None:
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = RGBColor(*fill_rgb)
                    else:
                        shape.fill.background()

                    # Stroke
                    if stroke_color is not None:
                        shape.line.color.rgb = RGBColor(*stroke_rgb)
                        shape.line.width = Pt(max(line_width, 0.25))
                    else:
                        shape.line.fill.background()

                    continue

                # 2) Straight line items.
                for item in items:
                    if not item:
                        continue

                    op = item[0]

                    if op == "l" and len(item) >= 3:
                        p1 = item[1]
                        p2 = item[2]

                        connector = slide.shapes.add_connector(
                            MSO_CONNECTOR.STRAIGHT,
                            Emu(map_x(p1.x)),
                            Emu(map_y(p1.y)),
                            Emu(map_x(p2.x)),
                            Emu(map_y(p2.y)),
                        )

                        if stroke_color is not None:
                            connector.line.color.rgb = RGBColor(*stroke_rgb)

                        connector.line.width = Pt(max(line_width, 0.25))

            except Exception:
                # Skip unsupported drawing safely.
                continue


    def _add_pdf_image_blocks_to_slide(
        self,
        slide,
        page: fitz.Page,
        image_left: int,
        image_top: int,
        image_width: int,
        image_height: int,
    ) -> None:
        """
        Add embedded PDF image blocks as PowerPoint image objects.

        This is used mainly for editable/text-only mode, where the full-page
        rendered image is not inserted. It keeps pictures/logos/screenshots
        visible while text remains editable.
        """
        try:
            from pptx.util import Emu
        except ImportError as exc:
            raise RuntimeError("python-pptx is required for image block export.") from exc

        page_rect = page.rect
        page_width = max(float(page_rect.width), 1.0)
        page_height = max(float(page_rect.height), 1.0)

        scale_x = image_width / page_width
        scale_y = image_height / page_height

        try:
            data = page.get_text("dict")
        except Exception:
            return

        for block in data.get("blocks", []):
            if block.get("type") != 1:
                continue

            bbox = block.get("bbox")
            image_bytes = block.get("image")
            image_ext = block.get("ext", "png")

            if not bbox or len(bbox) != 4 or not image_bytes:
                continue

            try:
                x0, y0, x1, y1 = [float(v) for v in bbox]

                left = image_left + int(x0 * scale_x)
                top = image_top + int(y0 * scale_y)
                width = max(int((x1 - x0) * scale_x), 1)
                height = max(int((y1 - y0) * scale_y), 1)

                stream = BytesIO(image_bytes)
                stream.seek(0)

                slide.shapes.add_picture(
                    stream,
                    Emu(left),
                    Emu(top),
                    width=Emu(width),
                    height=Emu(height),
                )

            except Exception:
                # Some PDF image blocks can be masks or unsupported streams.
                # Skip them rather than failing the whole conversion.
                continue


    def _add_editable_text_layer_to_slide(
        self,
        slide,
        page: fitz.Page,
        image_left: int,
        image_top: int,
        image_width: int,
        image_height: int,
    ) -> None:
        """
        Add PDF text as editable PowerPoint text boxes.

        v2 quality strategy:
        - Create one text box per visual line, not per span.
        - Estimate missing spaces between spans.
        - Use CJK-friendly font fallback for Korean/Japanese/Chinese text.
        """
        try:
            from pptx.util import Emu, Pt
            from pptx.dml.color import RGBColor
        except ImportError as exc:
            raise RuntimeError("python-pptx is required for editable text layer.") from exc

        page_rect = page.rect
        page_width = max(float(page_rect.width), 1.0)
        page_height = max(float(page_rect.height), 1.0)

        scale_x = image_width / page_width
        scale_y = image_height / page_height
        emu_per_pt = EMU_PER_INCH / PT_PER_INCH

        data = page.get_text("dict")

        for block in data.get("blocks", []):
            if block.get("type") != 0:
                continue

            for line in block.get("lines", []):
                spans = [
                    span for span in line.get("spans", [])
                    if _safe_text(span.get("text", ""))
                ]

                if not spans:
                    continue

                line_text = _remove_leaked_brand_tokens(_aggressive_korean_spacing(_final_korean_spacing_cleanup(_apply_korean_spacing_heuristics(_join_spans_with_estimated_spaces(spans)))))

                line_text = _fix_broken_pdf_symbols(line_text)

                if not line_text:
                    continue

                # Use union bbox for the full line.
                x0 = min(float(span.get("bbox", [0, 0, 0, 0])[0]) for span in spans)
                y0 = min(float(span.get("bbox", [0, 0, 0, 0])[1]) for span in spans)
                x1 = max(float(span.get("bbox", [0, 0, 0, 0])[2]) for span in spans)
                y1 = max(float(span.get("bbox", [0, 0, 0, 0])[3]) for span in spans)

                left = image_left + int(x0 * scale_x)
                top = image_top + int(y0 * scale_y)

                # Add extra width to avoid wrapping/truncation in PowerPoint.
                width = max(int((x1 - x0) * scale_x) + 24, 20)
                height = max(int((y1 - y0) * scale_y) + 8, 12)

                textbox = slide.shapes.add_textbox(
                    Emu(left),
                    Emu(top),
                    Emu(width),
                    Emu(height),
                )

                text_frame = textbox.text_frame
                text_frame.clear()
                text_frame.margin_left = 0
                text_frame.margin_right = 0
                text_frame.margin_top = 0
                text_frame.margin_bottom = 0
                text_frame.word_wrap = False
                text_frame.auto_size = None

                paragraph = text_frame.paragraphs[0]
                paragraph.space_after = Pt(0)
                paragraph.space_before = Pt(0)

                run = paragraph.add_run()
                run.text = line_text

                font = run.font

                # Representative span: largest text in this line.
                base_span = max(spans, key=lambda s: float(s.get("size", 10.0) or 10.0))

                try:
                    visual_scale_y = scale_y / emu_per_pt
                    base_size = float(base_span.get("size", 10.0) or 10.0)

                    # Slightly reduce to avoid visually oversized overlay.
                    font_size_pt = base_size * visual_scale_y * 0.92
                    font.size = Pt(max(font_size_pt, 1.0))
                except Exception:
                    font.size = Pt(10)

                font_name = base_span.get("font")

                if _contains_cjk(line_text):
                    font.name = "Malgun Gothic"
                elif font_name:
                    font.name = str(font_name).replace("+", "")

                r, g, b = _pdf_color_int_to_rgb(base_span.get("color", 0))
                font.color.rgb = RGBColor(r, g, b)

                flags = int(base_span.get("flags", 0) or 0)

                if flags & 16:
                    font.bold = True
                if flags & 2:
                    font.italic = True

    def export_pdf_to_word(
        self,
        source_pdf_path: str,
        output_docx_path: str,
        options: OfficeExportOptions | None = None,
    ) -> None:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is not installed. Run: pip install python-docx"
            ) from exc

        options = options or OfficeExportOptions()
        source_path = Path(source_pdf_path)
        output_path = Path(output_docx_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        pdf = fitz.open(str(source_path))

        try:
            page_indexes = options.page_range or list(range(len(pdf)))

            if not page_indexes:
                raise RuntimeError("No pages selected for Word export.")

            docx = Document()
            docx.add_heading(source_path.name, level=1)

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                text = page.get_text("text").strip()

                docx.add_heading(f"Page {page_index + 1}", level=2)

                if text:
                    for paragraph in re.split(r"\n\s*\n", text):
                        paragraph = paragraph.strip()
                        if paragraph:
                            docx.add_paragraph(paragraph)
                else:
                    docx.add_paragraph("[No extractable text on this page]")

                if options.include_page_breaks and order != len(page_indexes) - 1:
                    docx.add_page_break()

            docx.save(str(output_path))

        finally:
            pdf.close()

    def export_pdf_to_excel(
        self,
        source_pdf_path: str,
        output_xlsx_path: str,
        options: OfficeExportOptions | None = None,
    ) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl is not installed. Run: pip install openpyxl"
            ) from exc

        options = options or OfficeExportOptions()
        source_path = Path(source_pdf_path)
        output_path = Path(output_xlsx_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        pdf = fitz.open(str(source_path))

        try:
            page_indexes = options.page_range or list(range(len(pdf)))

            if not page_indexes:
                raise RuntimeError("No pages selected for Excel export.")

            wb = Workbook()
            ws = wb.active
            ws.title = "PDF Text"

            headers = ["Page", "Line No", "Text"]
            ws.append(headers)

            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(vertical="top")

            row_no = 2

            for page_index in page_indexes:
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                lines = page.get_text("text").splitlines()

                line_no = 1

                for line in lines:
                    line = line.strip()

                    if not line:
                        continue

                    ws.cell(row=row_no, column=1, value=page_index + 1)
                    ws.cell(row=row_no, column=2, value=line_no)
                    ws.cell(row=row_no, column=3, value=line)

                    row_no += 1
                    line_no += 1

            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 10
            ws.column_dimensions["C"].width = 100

            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            wb.save(str(output_path))

        finally:
            pdf.close()



# --- drawing quality boost override v1 ---
def _ppt_color_from_pdf_color(color_value, default=(0, 0, 0)):
    try:
        if color_value is None:
            return default

        if isinstance(color_value, (list, tuple)) and len(color_value) >= 3:
            return (
                int(max(0.0, min(1.0, float(color_value[0]))) * 255),
                int(max(0.0, min(1.0, float(color_value[1]))) * 255),
                int(max(0.0, min(1.0, float(color_value[2]))) * 255),
            )
    except Exception:
        pass

    return default


def _drawing_rect_to_values(rect):
    return float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)


def _boosted_add_pdf_drawings_to_slide(
    self,
    slide,
    page: fitz.Page,
    image_left: int,
    image_top: int,
    image_width: int,
    image_height: int,
    max_drawing_objects: int = 1200,
) -> None:
    """
    Improved simple vector drawing conversion.

    Supported:
    - lines
    - thin rectangles as lines
    - filled/stroked rectangles
    - simple rectangle path items

    Complex curves/freeforms are intentionally skipped for stability.
    """
    if max_drawing_objects <= 0:
        return

    try:
        from pptx.util import Emu, Pt
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.dml.color import RGBColor
    except Exception:
        return

    page_rect = page.rect
    page_width = max(float(page_rect.width), 1.0)
    page_height = max(float(page_rect.height), 1.0)

    scale_x = image_width / page_width
    scale_y = image_height / page_height
    scale_avg = max((scale_x + scale_y) / 2.0, 1.0)

    def map_x(x: float) -> int:
        return image_left + int(float(x) * scale_x)

    def map_y(y: float) -> int:
        return image_top + int(float(y) * scale_y)

    def map_w(w: float) -> int:
        return max(int(abs(float(w)) * scale_x), 1)

    def map_h(h: float) -> int:
        return max(int(abs(float(h)) * scale_y), 1)

    def set_line_style(shape, color, width):
        try:
            r, g, b = _ppt_color_from_pdf_color(color, default=(0, 0, 0))
            shape.line.color.rgb = RGBColor(r, g, b)

            # PDF line width is in points. Slightly boost very thin lines.
            width_pt = max(float(width or 0.5) * 1.15, 0.35)
            shape.line.width = Pt(width_pt)
        except Exception:
            pass

    def set_fill_style(shape, fill_color):
        try:
            if fill_color is None:
                shape.fill.background()
                return

            r, g, b = _ppt_color_from_pdf_color(fill_color, default=(255, 255, 255))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(r, g, b)
        except Exception:
            try:
                shape.fill.background()
            except Exception:
                pass

    def add_line(x0, y0, x1, y1, color, width):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Emu(map_x(x0)),
            Emu(map_y(y0)),
            Emu(map_x(x1)),
            Emu(map_y(y1)),
        )
        set_line_style(connector, color, width)
        return connector

    def add_rectangle(x0, y0, x1, y1, stroke_color, fill_color, width):
        left = map_x(min(x0, x1))
        top = map_y(min(y0, y1))
        w = map_w(x1 - x0)
        h = map_h(y1 - y0)

        # Skip tiny noise.
        if w <= 1 and h <= 1:
            return None

        # Thin rectangle fallback to line.
        if w <= 2 and h > 2:
            return add_line(x0, y0, x0, y1, stroke_color or fill_color, width)

        if h <= 2 and w > 2:
            return add_line(x0, y0, x1, y0, stroke_color or fill_color, width)

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(left),
            Emu(top),
            Emu(w),
            Emu(h),
        )

        set_fill_style(shape, fill_color)

        if stroke_color is not None:
            set_line_style(shape, stroke_color, width)
        else:
            try:
                shape.line.fill.background()
            except Exception:
                pass

        return shape

    try:
        drawings = page.get_drawings()
    except Exception:
        return

    if not drawings:
        return

    converted = 0

    for drawing in drawings:
        if converted >= max_drawing_objects:
            break

        try:
            stroke_color = drawing.get("color")
            fill_color = drawing.get("fill")
            line_width = float(drawing.get("width", 0.5) or 0.5)
            rect = drawing.get("rect")
            items = drawing.get("items", [])

            # Skip extremely complex paths to avoid freezing.
            if len(items) > 40:
                continue

            # Rectangle path items have priority.
            handled = False

            for item in items:
                if not item:
                    continue

                op = item[0]

                # Rectangle item: ("re", Rect, ...)
                if op == "re" and len(item) >= 2:
                    r = item[1]
                    x0, y0, x1, y1 = _drawing_rect_to_values(r)
                    shape = add_rectangle(
                        x0,
                        y0,
                        x1,
                        y1,
                        stroke_color,
                        fill_color,
                        line_width,
                    )
                    if shape is not None:
                        converted += 1
                    handled = True

                    if converted >= max_drawing_objects:
                        break

                # Straight line item: ("l", Point, Point)
                elif op == "l" and len(item) >= 3:
                    p1 = item[1]
                    p2 = item[2]

                    # Skip tiny noise.
                    if abs(float(p1.x) - float(p2.x)) < 0.2 and abs(float(p1.y) - float(p2.y)) < 0.2:
                        continue

                    add_line(
                        float(p1.x),
                        float(p1.y),
                        float(p2.x),
                        float(p2.y),
                        stroke_color,
                        line_width,
                    )
                    converted += 1
                    handled = True

                    if converted >= max_drawing_objects:
                        break

            if handled:
                continue

            # Fallback: if drawing has fill and a bounding rect, convert as rectangle.
            if rect is not None and fill_color is not None:
                x0, y0, x1, y1 = _drawing_rect_to_values(rect)

                # Ignore huge full-page fills; they often cover everything.
                rect_area = abs((x1 - x0) * (y1 - y0))
                page_area = page_width * page_height

                if rect_area < page_area * 0.95:
                    shape = add_rectangle(
                        x0,
                        y0,
                        x1,
                        y1,
                        stroke_color,
                        fill_color,
                        line_width,
                    )
                    if shape is not None:
                        converted += 1

        except Exception:
            continue


# Override class method safely.
try:
    OfficeExporter._add_pdf_drawings_to_slide = _boosted_add_pdf_drawings_to_slide
except NameError:
    pass



# --- word based ppt text engine override v1 ---
def _group_pdf_words_into_lines(page: fitz.Page) -> list[dict]:
    """
    Build visual lines from PyMuPDF words.

    page.get_text("words") returns:
    (x0, y0, x1, y1, word, block_no, line_no, word_no)

    This is generally better for spacing than span-level extraction.
    """
    try:
        words = page.get_text("words")
    except Exception:
        return []

    if not words:
        return []

    grouped: dict[tuple[int, int], list[tuple]] = {}

    for word in words:
        if len(word) < 8:
            continue

        x0, y0, x1, y1, text, block_no, line_no, word_no = word[:8]
        text = _fix_broken_pdf_symbols(str(text or "")).strip()

        if not text:
            continue

        key = (int(block_no), int(line_no))
        grouped.setdefault(key, []).append((float(x0), float(y0), float(x1), float(y1), text, int(word_no)))

    lines: list[dict] = []

    for key in sorted(grouped.keys()):
        items = grouped[key]
        items.sort(key=lambda item: (item[5], item[0]))

        x0 = min(item[0] for item in items)
        y0 = min(item[1] for item in items)
        x1 = max(item[2] for item in items)
        y1 = max(item[3] for item in items)

        parts: list[str] = []
        prev_x1 = None
        prev_text = ""

        heights = [max(item[3] - item[1], 1.0) for item in items]
        avg_height = sum(heights) / max(len(heights), 1)

        for item in items:
            wx0, wy0, wx1, wy1, word_text, _word_no = item

            if prev_x1 is not None:
                gap = wx0 - prev_x1

                # Estimate visual spacing.
                # For word extraction, most separate words are already separated,
                # but some Korean/ASCII boundaries still need help.
                if gap > max(avg_height * 0.12, 1.2):
                    if parts and not parts[-1].endswith(" "):
                        parts.append(" ")

            parts.append(word_text)
            prev_x1 = wx1
            prev_text = word_text

        line_text = "".join(parts)
        line_text = _normalize_line_text(line_text)
        line_text = _apply_korean_spacing_heuristics(line_text)
        line_text = _final_korean_spacing_cleanup(line_text)
        line_text = _aggressive_korean_spacing(line_text)
        line_text = _remove_leaked_brand_tokens(line_text)
        line_text = _fix_broken_pdf_symbols(line_text)

        if not line_text:
            continue

        lines.append(
            {
                "text": line_text,
                "bbox": (x0, y0, x1, y1),
                "avg_height": avg_height,
            }
        )

    return lines


def _word_based_add_editable_text_layer_to_slide(
    self,
    slide,
    page: fitz.Page,
    image_left: int,
    image_top: int,
    image_width: int,
    image_height: int,
) -> None:
    """
    Add editable PowerPoint text boxes using PDF word extraction.

    Benefits:
    - Better spacing than span-based mode
    - Fewer text boxes
    - More natural line editing
    """
    try:
        from pptx.util import Emu, Pt
        from pptx.dml.color import RGBColor
    except Exception as exc:
        raise RuntimeError("python-pptx is required for editable text layer.") from exc

    page_rect = page.rect
    page_width = max(float(page_rect.width), 1.0)
    page_height = max(float(page_rect.height), 1.0)

    scale_x = image_width / page_width
    scale_y = image_height / page_height
    emu_per_pt = EMU_PER_INCH / PT_PER_INCH

    lines = _group_pdf_words_into_lines(page)

    # Fallback to previous span method if word extraction fails.
    if not lines:
        try:
            return _old_span_add_editable_text_layer_to_slide(
                self,
                slide,
                page,
                image_left,
                image_top,
                image_width,
                image_height,
            )
        except Exception:
            return

    for line in lines:
        try:
            text = line["text"]
            x0, y0, x1, y1 = line["bbox"]
            avg_height = float(line.get("avg_height", max(y1 - y0, 10.0)))

            left = image_left + int(x0 * scale_x)
            top = image_top + int(y0 * scale_y)
            width = max(int((x1 - x0) * scale_x) + 26, 20)
            height = max(int((y1 - y0) * scale_y) + 10, 12)

            textbox = slide.shapes.add_textbox(
                Emu(left),
                Emu(top),
                Emu(width),
                Emu(height),
            )

            # For no-overlap editable mode this background is usually harmless.
            # It also helps when a user later manually places it above images.
            try:
                from pptx.dml.color import RGBColor
                textbox.fill.background()
                textbox.line.fill.background()
            except Exception:
                pass

            text_frame = textbox.text_frame
            text_frame.clear()
            text_frame.margin_left = 0
            text_frame.margin_right = 0
            text_frame.margin_top = 0
            text_frame.margin_bottom = 0
            text_frame.word_wrap = False

            paragraph = text_frame.paragraphs[0]
            paragraph.space_after = Pt(0)
            paragraph.space_before = Pt(0)

            run = paragraph.add_run()
            run.text = text

            font = run.font

            # Estimate font size from word bbox height.
            try:
                visual_scale_y = scale_y / emu_per_pt
                font_size_pt = avg_height * visual_scale_y * 0.78
                font.size = Pt(max(font_size_pt, 3.0))
            except Exception:
                font.size = Pt(10)

            if _contains_cjk(text):
                font.name = "Malgun Gothic"
            else:
                font.name = "Arial"

            font.color.rgb = RGBColor(0, 0, 0)

        except Exception:
            continue


# Preserve old span-based method as fallback once.
try:
    if not hasattr(OfficeExporter, "_old_span_add_editable_text_layer_to_slide"):
        OfficeExporter._old_span_add_editable_text_layer_to_slide = OfficeExporter._add_editable_text_layer_to_slide
    _old_span_add_editable_text_layer_to_slide = OfficeExporter._old_span_add_editable_text_layer_to_slide
    OfficeExporter._add_editable_text_layer_to_slide = _word_based_add_editable_text_layer_to_slide
except Exception:
    pass



# --- complex drawing fallback override v2 ---
def _ppt_color_from_pdf_color_v2(color_value, default=(0, 0, 0)):
    try:
        if color_value is None:
            return default
        if isinstance(color_value, (list, tuple)) and len(color_value) >= 3:
            return (
                int(max(0.0, min(1.0, float(color_value[0]))) * 255),
                int(max(0.0, min(1.0, float(color_value[1]))) * 255),
                int(max(0.0, min(1.0, float(color_value[2]))) * 255),
            )
    except Exception:
        pass
    return default


def _rect_values_v2(rect):
    return float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)


def _boost_rect_v2(rect: fitz.Rect, amount: float, page_rect: fitz.Rect) -> fitz.Rect:
    return fitz.Rect(
        max(page_rect.x0, rect.x0 - amount),
        max(page_rect.y0, rect.y0 - amount),
        min(page_rect.x1, rect.x1 + amount),
        min(page_rect.y1, rect.y1 + amount),
    )


def _add_clip_image_to_slide_v2(
    slide,
    page: fitz.Page,
    clip: fitz.Rect,
    image_left: int,
    image_top: int,
    image_width: int,
    image_height: int,
    scale_x: float,
    scale_y: float,
    dpi: int = 220,
) -> bool:
    """
    Render a small PDF clip as PNG and insert into slide.
    Used for complex vector icons/symbols that are hard to recreate as PPT shapes.
    """
    try:
        from pptx.util import Emu
    except Exception:
        return False

    try:
        if clip.width <= 0 or clip.height <= 0:
            return False

        zoom = max(float(dpi), 96.0) / 72.0
        pix = page.get_pixmap(
            matrix=fitz.Matrix(zoom, zoom),
            clip=clip,
            alpha=True,
            colorspace=fitz.csRGB,
        )

        png_bytes = pix.tobytes("png")
        stream = BytesIO(png_bytes)
        stream.seek(0)

        left = image_left + int(clip.x0 * scale_x)
        top = image_top + int(clip.y0 * scale_y)
        width = max(int(clip.width * scale_x), 1)
        height = max(int(clip.height * scale_y), 1)

        slide.shapes.add_picture(
            stream,
            Emu(left),
            Emu(top),
            width=Emu(width),
            height=Emu(height),
        )
        return True

    except Exception:
        return False


def _complex_fallback_add_pdf_drawings_to_slide(
    self,
    slide,
    page: fitz.Page,
    image_left: int,
    image_top: int,
    image_width: int,
    image_height: int,
    max_drawing_objects: int = 1200,
) -> None:
    """
    Drawing engine v2:
    - Convert simple lines/rectangles as PPT shapes.
    - Convert complex small vector paths as clipped images.
    - Avoid full-page fallback to prevent text overlap.
    """
    if max_drawing_objects <= 0:
        return

    try:
        from pptx.util import Emu, Pt
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.dml.color import RGBColor
    except Exception:
        return

    page_rect = page.rect
    page_width = max(float(page_rect.width), 1.0)
    page_height = max(float(page_rect.height), 1.0)
    page_area = page_width * page_height

    scale_x = image_width / page_width
    scale_y = image_height / page_height

    def map_x(x: float) -> int:
        return image_left + int(float(x) * scale_x)

    def map_y(y: float) -> int:
        return image_top + int(float(y) * scale_y)

    def map_w(w: float) -> int:
        return max(int(abs(float(w)) * scale_x), 1)

    def map_h(h: float) -> int:
        return max(int(abs(float(h)) * scale_y), 1)

    def set_line_style(shape, color, width):
        try:
            r, g, b = _ppt_color_from_pdf_color_v2(color, default=(0, 0, 0))
            shape.line.color.rgb = RGBColor(r, g, b)
            shape.line.width = Pt(max(float(width or 0.5) * 1.2, 0.35))
        except Exception:
            pass

    def set_fill_style(shape, fill_color):
        try:
            if fill_color is None:
                shape.fill.background()
                return
            r, g, b = _ppt_color_from_pdf_color_v2(fill_color, default=(255, 255, 255))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(r, g, b)
        except Exception:
            try:
                shape.fill.background()
            except Exception:
                pass

    def add_line(x0, y0, x1, y1, color, width):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Emu(map_x(x0)),
            Emu(map_y(y0)),
            Emu(map_x(x1)),
            Emu(map_y(y1)),
        )
        set_line_style(connector, color, width)
        return connector

    def add_rectangle(x0, y0, x1, y1, stroke_color, fill_color, width):
        left = map_x(min(x0, x1))
        top = map_y(min(y0, y1))
        w = map_w(x1 - x0)
        h = map_h(y1 - y0)

        if w <= 1 and h <= 1:
            return None

        # Thin rectangles are usually borders/lines.
        if w <= 2 and h > 2:
            return add_line(x0, y0, x0, y1, stroke_color or fill_color, width)
        if h <= 2 and w > 2:
            return add_line(x0, y0, x1, y0, stroke_color or fill_color, width)

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(left),
            Emu(top),
            Emu(w),
            Emu(h),
        )

        set_fill_style(shape, fill_color)

        if stroke_color is not None:
            set_line_style(shape, stroke_color, width)
        else:
            try:
                shape.line.fill.background()
            except Exception:
                pass

        return shape

    try:
        drawings = page.get_drawings()
    except Exception:
        return

    if not drawings:
        return

    converted = 0
    fallback_images = 0
    max_fallback_images = 180

    for drawing in drawings:
        if converted >= max_drawing_objects:
            break

        try:
            stroke_color = drawing.get("color")
            fill_color = drawing.get("fill")
            line_width = float(drawing.get("width", 0.5) or 0.5)
            rect = drawing.get("rect")
            items = drawing.get("items", [])

            if rect is None:
                continue

            clip = fitz.Rect(rect)
            clip_area = max(clip.width * clip.height, 0.0)

            # Skip full-page or huge backgrounds. They would cause text overlap.
            if clip_area <= 0 or clip_area > page_area * 0.35:
                continue

            # Skip microscopic noise.
            if clip.width < 0.8 or clip.height < 0.8:
                continue

            simple_count = 0
            complex_count = 0
            handled = False

            for item in items:
                if not item:
                    continue

                op = item[0]

                if op == "re" and len(item) >= 2:
                    r = item[1]
                    x0, y0, x1, y1 = _rect_values_v2(r)
                    shape = add_rectangle(
                        x0,
                        y0,
                        x1,
                        y1,
                        stroke_color,
                        fill_color,
                        line_width,
                    )
                    if shape is not None:
                        converted += 1
                        simple_count += 1
                        handled = True

                elif op == "l" and len(item) >= 3:
                    p1 = item[1]
                    p2 = item[2]

                    if abs(float(p1.x) - float(p2.x)) < 0.2 and abs(float(p1.y) - float(p2.y)) < 0.2:
                        continue

                    add_line(
                        float(p1.x),
                        float(p1.y),
                        float(p2.x),
                        float(p2.y),
                        stroke_color,
                        line_width,
                    )
                    converted += 1
                    simple_count += 1
                    handled = True

                else:
                    # curve/path/quad/etc.
                    complex_count += 1

                if converted >= max_drawing_objects:
                    break

            if handled:
                continue

            # Filled bounding rectangle fallback.
            if fill_color is not None and len(items) <= 4:
                x0, y0, x1, y1 = _rect_values_v2(clip)
                shape = add_rectangle(
                    x0,
                    y0,
                    x1,
                    y1,
                    stroke_color,
                    fill_color,
                    line_width,
                )
                if shape is not None:
                    converted += 1
                    continue

            # Complex vector fallback as clipped image.
            # Use only small/medium non-text-like areas to reduce overlap risk.
            if fallback_images < max_fallback_images:
                is_small_icon = clip_area < page_area * 0.04
                is_medium_graphic = clip_area < page_area * 0.12 and len(items) >= 3

                if complex_count > 0 and (is_small_icon or is_medium_graphic):
                    boosted = _boost_rect_v2(clip, 1.0, page_rect)

                    ok = _add_clip_image_to_slide_v2(
                        slide=slide,
                        page=page,
                        clip=boosted,
                        image_left=image_left,
                        image_top=image_top,
                        image_width=image_width,
                        image_height=image_height,
                        scale_x=scale_x,
                        scale_y=scale_y,
                        dpi=220,
                    )

                    if ok:
                        fallback_images += 1
                        converted += 1

        except Exception:
            continue


try:
    OfficeExporter._add_pdf_drawings_to_slide = _complex_fallback_add_pdf_drawings_to_slide
except NameError:
    pass



# --- powerpoint auto mode override v1 ---
def _analyze_pdf_page_complexity_for_ppt(page: fitz.Page) -> dict:
    """
    Decide whether a page is suitable for editable conversion.

    Complex technical drawings, board layouts, schematics, dense vector pages
    are better exported as a single image slide.
    """
    result = {
        "drawings": 0,
        "drawing_items": 0,
        "words": 0,
        "images": 0,
        "small_drawings": 0,
        "complex": False,
        "reason": "",
    }

    try:
        drawings = page.get_drawings()
        result["drawings"] = len(drawings)

        page_area = max(float(page.rect.width * page.rect.height), 1.0)

        item_count = 0
        small_count = 0

        for d in drawings:
            items = d.get("items", [])
            item_count += len(items)

            rect = d.get("rect")
            if rect is not None:
                area = max(float(rect.width * rect.height), 0.0)
                if area > 0 and area < page_area * 0.002:
                    small_count += 1

        result["drawing_items"] = item_count
        result["small_drawings"] = small_count
    except Exception:
        pass

    try:
        result["words"] = len(page.get_text("words"))
    except Exception:
        pass

    try:
        result["images"] = len(page.get_images(full=True))
    except Exception:
        pass

    # Heuristic:
    # Mainboard layouts / technical diagrams have many vector drawings and small items.
    if result["drawings"] >= 250:
        result["complex"] = True
        result["reason"] = f"many drawings ({result['drawings']})"
    elif result["drawing_items"] >= 900:
        result["complex"] = True
        result["reason"] = f"many drawing items ({result['drawing_items']})"
    elif result["small_drawings"] >= 120:
        result["complex"] = True
        result["reason"] = f"many small drawings ({result['small_drawings']})"
    elif result["words"] >= 250 and result["drawings"] >= 120:
        result["complex"] = True
        result["reason"] = "dense mixed text/vector page"
    else:
        result["reason"] = "editable-friendly page"

    return result


def _auto_mode_export_pdf_to_powerpoint(
    self,
    source_pdf_path: str,
    output_pptx_path: str,
    options: OfficeExportOptions | None = None,
    progress_callback=None,
) -> None:
    """
    PowerPoint export with page-level auto fallback.

    Modes:
    - image: all pages as image slides
    - text_only: editable text/images/drawings
    - auto: complex pages as image, simpler pages as editable
    """
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is not installed. Run: pip install python-pptx Pillow"
        ) from exc

    options = options or OfficeExportOptions()
    options.auto_mode_report = []
    source_path = Path(source_pdf_path)
    output_path = Path(output_pptx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = fitz.open(str(source_path))

    try:
        page_indexes = options.page_range or list(range(len(doc)))

        if not page_indexes:
            raise RuntimeError("No pages selected for PowerPoint export.")

        first_page = doc.load_page(page_indexes[0])
        first_rect = first_page.rect

        prs = Presentation()
        prs.slide_width = Emu(int((first_rect.width / PT_PER_INCH) * EMU_PER_INCH))
        prs.slide_height = Emu(int((first_rect.height / PT_PER_INCH) * EMU_PER_INCH))

        blank_layout = prs.slide_layouts[6]
        zoom = max(options.dpi, 72) / 72.0
        total_pages = len(page_indexes)

        import tempfile

        with tempfile.TemporaryDirectory(prefix="pdf_pptx_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            for output_no, page_index in enumerate(page_indexes, start=1):
                if progress_callback:
                    if progress_callback(
                        output_no - 1,
                        total_pages,
                        f"페이지 {output_no}/{total_pages} 분석 중..."
                    ) is False:
                        raise RuntimeError("PowerPoint conversion was cancelled.")

                if page_index < 0 or page_index >= len(doc):
                    continue

                page = doc.load_page(page_index)
                page_rect = page.rect

                requested_mode = getattr(options, "powerpoint_mode", "image")
                effective_mode = requested_mode

                complexity = None
                if requested_mode == "auto":
                    complexity = _analyze_pdf_page_complexity_for_ppt(page)
                    effective_mode = "image" if complexity.get("complex") else "text_only"

                    try:
                        report_item = {
                            "page": page_index + 1,
                            "mode": effective_mode,
                            "reason": complexity.get("reason", ""),
                            "drawings": complexity.get("drawings", 0),
                            "drawing_items": complexity.get("drawing_items", 0),
                            "small_drawings": complexity.get("small_drawings", 0),
                            "words": complexity.get("words", 0),
                            "images": complexity.get("images", 0),
                        }
                        options.auto_mode_report.append(report_item)

                        print(
                            f"[AUTO] Page {page_index + 1}: {effective_mode} - "
                            f"{report_item['reason']} "
                            f"(drawings={report_item['drawings']}, "
                            f"items={report_item['drawing_items']}, "
                            f"small={report_item['small_drawings']}, "
                            f"words={report_item['words']}, "
                            f"images={report_item['images']})"
                        )
                    except Exception:
                        pass

                    try:
                        report_item = {
                            "page": page_index + 1,
                            "mode": effective_mode,
                            "reason": complexity.get("reason", ""),
                            "drawings": complexity.get("drawings", 0),
                            "drawing_items": complexity.get("drawing_items", 0),
                            "small_drawings": complexity.get("small_drawings", 0),
                            "words": complexity.get("words", 0),
                            "images": complexity.get("images", 0),
                        }
                        options.auto_mode_report.append(report_item)

                        print(
                            f"[AUTO] Page {page_index + 1}: {effective_mode} - "
                            f"{report_item['reason']} "
                            f"(drawings={report_item['drawings']}, "
                            f"items={report_item['drawing_items']}, "
                            f"small={report_item['small_drawings']}, "
                            f"words={report_item['words']}, "
                            f"images={report_item['images']})"
                        )
                    except Exception:
                        pass

                if progress_callback:
                    mode_text = "원본 이미지" if effective_mode == "image" else "편집 가능"
                    reason = f" ({complexity.get('reason')})" if complexity else ""
                    if progress_callback(
                        output_no - 1,
                        total_pages,
                        f"페이지 {output_no}/{total_pages} 처리 중: {mode_text}{reason}"
                    ) is False:
                        raise RuntimeError("PowerPoint conversion was cancelled.")

                pix = page.get_pixmap(
                    matrix=fitz.Matrix(zoom, zoom),
                    alpha=False,
                    colorspace=fitz.csRGB,
                )

                image_path = temp_dir_path / f"page_{output_no:04d}.png"
                pix.save(str(image_path))

                slide = prs.slides.add_slide(blank_layout)

                slide_w = prs.slide_width
                slide_h = prs.slide_height

                page_ratio = page_rect.width / page_rect.height
                slide_ratio = int(slide_w) / int(slide_h)

                if getattr(options, "fit_mode", "fit") == "fill":
                    if page_ratio > slide_ratio:
                        image_height = slide_h
                        image_width = int(image_height * page_ratio)
                    else:
                        image_width = slide_w
                        image_height = int(image_width / page_ratio)
                else:
                    if page_ratio > slide_ratio:
                        image_width = slide_w
                        image_height = int(image_width / page_ratio)
                    else:
                        image_height = slide_h
                        image_width = int(image_height * page_ratio)

                image_left = int((int(slide_w) - int(image_width)) / 2)
                image_top = int((int(slide_h) - int(image_height)) / 2)

                if effective_mode == "image":
                    slide.shapes.add_picture(
                        str(image_path),
                        image_left,
                        image_top,
                        width=image_width,
                        height=image_height,
                    )
                else:
                    # Editable mode: no full-page background image to avoid text overlap.
                    # Add vector/image objects first, then editable text.
                    if progress_callback:
                        if progress_callback(
                            output_no - 1,
                            total_pages,
                            f"페이지 {output_no}/{total_pages} 이미지/도형 변환 중..."
                        ) is False:
                            raise RuntimeError("PowerPoint conversion was cancelled.")

                    try:
                        self._add_pdf_drawings_to_slide(
                            slide=slide,
                            page=page,
                            image_left=image_left,
                            image_top=image_top,
                            image_width=int(image_width),
                            image_height=int(image_height),
                            max_drawing_objects=getattr(options, "max_drawing_objects", 1200),
                        )
                    except Exception:
                        pass

                    try:
                        self._add_pdf_image_blocks_to_slide(
                            slide=slide,
                            page=page,
                            image_left=image_left,
                            image_top=image_top,
                            image_width=int(image_width),
                            image_height=int(image_height),
                        )
                    except Exception:
                        pass

                    if progress_callback:
                        if progress_callback(
                            output_no - 1,
                            total_pages,
                            f"페이지 {output_no}/{total_pages} 텍스트 변환 중..."
                        ) is False:
                            raise RuntimeError("PowerPoint conversion was cancelled.")

                    self._current_text_box_fill_transparency = 100.0

                    try:
                        self._add_editable_text_layer_to_slide(
                            slide=slide,
                            page=page,
                            image_left=image_left,
                            image_top=image_top,
                            image_width=int(image_width),
                            image_height=int(image_height),
                        )
                    except Exception:
                        pass

                if progress_callback:
                    if progress_callback(
                        output_no,
                        total_pages,
                        f"페이지 {output_no}/{total_pages} 완료"
                    ) is False:
                        raise RuntimeError("PowerPoint conversion was cancelled.")

        prs.save(str(output_path))

    finally:
        doc.close()


try:
    OfficeExporter.export_pdf_to_powerpoint = _auto_mode_export_pdf_to_powerpoint
except NameError:
    pass



# --- improved auto complexity override v2 ---
def _analyze_pdf_page_complexity_for_ppt(page: fitz.Page) -> dict:
    """
    Decide whether a page should be converted as editable objects
    or preserved as a full-page image.

    Strongly fallback to image for:
    - technical drawings
    - PCB/mainboard layouts
    - dense vector diagrams
    - many tiny labels and objects
    """
    result = {
        "drawings": 0,
        "drawing_items": 0,
        "words": 0,
        "images": 0,
        "small_drawings": 0,
        "tiny_words": 0,
        "avg_word_height": 0.0,
        "complex": False,
        "reason": "",
    }

    page_area = max(float(page.rect.width * page.rect.height), 1.0)

    try:
        drawings = page.get_drawings()
        result["drawings"] = len(drawings)

        item_count = 0
        small_count = 0

        for d in drawings:
            items = d.get("items", [])
            item_count += len(items)

            rect = d.get("rect")
            if rect is not None:
                area = max(float(rect.width * rect.height), 0.0)
                if 0 < area < page_area * 0.002:
                    small_count += 1

        result["drawing_items"] = item_count
        result["small_drawings"] = small_count

    except Exception:
        pass

    try:
        words = page.get_text("words")
        result["words"] = len(words)

        heights = []
        tiny_count = 0

        for w in words:
            if len(w) < 5:
                continue

            y0 = float(w[1])
            y1 = float(w[3])
            h = max(y1 - y0, 0.0)
            heights.append(h)

            if h > 0 and h < 6.5:
                tiny_count += 1

        result["tiny_words"] = tiny_count
        result["avg_word_height"] = (
            sum(heights) / len(heights)
            if heights
            else 0.0
        )

    except Exception:
        pass

    try:
        result["images"] = len(page.get_images(full=True))
    except Exception:
        pass

    drawings = result["drawings"]
    items = result["drawing_items"]
    small = result["small_drawings"]
    words = result["words"]
    tiny_words = result["tiny_words"]
    avg_h = result["avg_word_height"]

    # Strong technical drawing indicators.
    if drawings >= 220:
        result["complex"] = True
        result["reason"] = f"many drawings ({drawings})"
    elif items >= 750:
        result["complex"] = True
        result["reason"] = f"many drawing items ({items})"
    elif small >= 90:
        result["complex"] = True
        result["reason"] = f"many small drawings ({small})"
    elif tiny_words >= 80 and drawings >= 80:
        result["complex"] = True
        result["reason"] = f"many tiny labels and drawings (tiny_words={tiny_words}, drawings={drawings})"
    elif words >= 180 and drawings >= 90:
        result["complex"] = True
        result["reason"] = f"dense mixed text/vector page (words={words}, drawings={drawings})"
    elif avg_h > 0 and avg_h < 6.0 and words >= 120:
        result["complex"] = True
        result["reason"] = f"many small labels (avg_word_height={avg_h:.1f})"
    else:
        result["reason"] = "editable-friendly page"

    return result



# --- word excel export improvement override v1 ---
def _improved_export_pdf_to_word(
    self,
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Improved Word export.

    Modes:
    - word_mode='text': text only
    - word_mode='image_text': page image + extracted text
    """
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    import tempfile

    options = options or OfficeExportOptions()
    word_mode = getattr(options, "word_mode", "text")

    source_path = Path(source_pdf_path)
    output_path = Path(output_docx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Word export.")

        docx = Document()
        docx.add_heading(source_path.name, level=1)

        with tempfile.TemporaryDirectory(prefix="pdf_docx_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                docx.add_heading(f"Page {page_index + 1}", level=2)

                if word_mode == "image_text":
                    try:
                        zoom = max(float(getattr(options, "dpi", 200)), 72.0) / 72.0
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(zoom, zoom),
                            alpha=False,
                            colorspace=fitz.csRGB,
                        )
                        image_path = temp_dir_path / f"page_{page_index + 1:04d}.png"
                        pix.save(str(image_path))

                        # Fit to typical Word page width.
                        docx.add_picture(str(image_path), width=Inches(6.5))
                    except Exception:
                        pass

                    docx.add_paragraph("")

                text = page.get_text("text").strip()
                text = _fix_broken_pdf_symbols(text)

                if text:
                    for paragraph in re.split(r"\n\s*\n", text):
                        paragraph = paragraph.strip()
                        if paragraph:
                            paragraph = _aggressive_korean_spacing(
                                _final_korean_spacing_cleanup(
                                    _apply_korean_spacing_heuristics(paragraph)
                                )
                            )
                            paragraph = _remove_leaked_brand_tokens(paragraph)
                            docx.add_paragraph(paragraph)
                else:
                    docx.add_paragraph("[No extractable text on this page]")

                if options.include_page_breaks and order != len(page_indexes) - 1:
                    docx.add_page_break()

        docx.save(str(output_path))

    finally:
        pdf.close()


def _group_words_to_table_rows(page: fitz.Page) -> list[list[str]]:
    """
    Coordinate-based table-like text extraction.

    This is heuristic:
    - group words by y position into rows
    - split cells by large x gaps
    """
    try:
        words = page.get_text("words")
    except Exception:
        return []

    if not words:
        return []

    items = []
    for w in words:
        if len(w) < 8:
            continue
        x0, y0, x1, y1, text, block_no, line_no, word_no = w[:8]
        text = _fix_broken_pdf_symbols(str(text or "")).strip()
        if not text:
            continue
        items.append(
            {
                "x0": float(x0),
                "y0": float(y0),
                "x1": float(x1),
                "y1": float(y1),
                "text": text,
                "h": max(float(y1) - float(y0), 1.0),
            }
        )

    if not items:
        return []

    items.sort(key=lambda it: (it["y0"], it["x0"]))

    rows: list[list[dict]] = []
    current: list[dict] = []
    current_y = None
    avg_h = sum(it["h"] for it in items) / len(items)
    y_threshold = max(avg_h * 0.55, 3.0)

    for it in items:
        if current_y is None:
            current = [it]
            current_y = it["y0"]
            continue

        if abs(it["y0"] - current_y) <= y_threshold:
            current.append(it)
            current_y = (current_y * (len(current) - 1) + it["y0"]) / len(current)
        else:
            rows.append(current)
            current = [it]
            current_y = it["y0"]

    if current:
        rows.append(current)

    table_rows: list[list[str]] = []

    for row in rows:
        row.sort(key=lambda it: it["x0"])

        if not row:
            continue

        heights = [it["h"] for it in row]
        row_h = sum(heights) / len(heights)
        gap_threshold = max(row_h * 1.2, 8.0)

        cells: list[str] = []
        current_cell_parts: list[str] = []
        prev_x1 = None

        for it in row:
            if prev_x1 is not None:
                gap = it["x0"] - prev_x1
                if gap > gap_threshold:
                    cell_text = " ".join(current_cell_parts).strip()
                    if cell_text:
                        cells.append(cell_text)
                    current_cell_parts = []

            current_cell_parts.append(it["text"])
            prev_x1 = it["x1"]

        if current_cell_parts:
            cells.append(" ".join(current_cell_parts).strip())

        cleaned_cells = []
        for cell in cells:
            cell = _normalize_line_text(cell)
            cell = _aggressive_korean_spacing(
                _final_korean_spacing_cleanup(
                    _apply_korean_spacing_heuristics(cell)
                )
            )
            cell = _remove_leaked_brand_tokens(cell)
            cleaned_cells.append(cell)

        if cleaned_cells:
            table_rows.append(cleaned_cells)

    return table_rows


def _improved_export_pdf_to_excel(
    self,
    source_pdf_path: str,
    output_xlsx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Improved Excel export.

    Modes:
    - excel_mode='lines': Page / Line No / Text
    - excel_mode='table': coordinate-based table estimation
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is not installed. Run: pip install openpyxl"
        ) from exc

    options = options or OfficeExportOptions()
    excel_mode = getattr(options, "excel_mode", "lines")

    source_path = Path(source_pdf_path)
    output_path = Path(output_xlsx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Excel export.")

        wb = Workbook()
        ws = wb.active
        ws.title = "PDF Text"

        if excel_mode == "table":
            current_row = 1

            for page_index in page_indexes:
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                rows = _group_words_to_table_rows(page)

                ws.cell(row=current_row, column=1, value=f"Page {page_index + 1}")
                ws.cell(row=current_row, column=1).font = Font(bold=True)
                current_row += 1

                for row in rows:
                    for col_index, value in enumerate(row, start=1):
                        ws.cell(row=current_row, column=col_index, value=value)
                    current_row += 1

                current_row += 1

            for col in range(1, 20):
                ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = 18

        else:
            headers = ["Page", "Line No", "Text"]
            ws.append(headers)

            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(vertical="top")

            row_no = 2

            for page_index in page_indexes:
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                lines = page.get_text("text").splitlines()

                line_no = 1

                for line in lines:
                    line = _fix_broken_pdf_symbols(line.strip())

                    if not line:
                        continue

                    line = _aggressive_korean_spacing(
                        _final_korean_spacing_cleanup(
                            _apply_korean_spacing_heuristics(line)
                        )
                    )
                    line = _remove_leaked_brand_tokens(line)

                    ws.cell(row=row_no, column=1, value=page_index + 1)
                    ws.cell(row=row_no, column=2, value=line_no)
                    ws.cell(row=row_no, column=3, value=line)

                    row_no += 1
                    line_no += 1

            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 10
            ws.column_dimensions["C"].width = 100

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        wb.save(str(output_path))

    finally:
        pdf.close()


try:
    OfficeExporter.export_pdf_to_word = _improved_export_pdf_to_word
    OfficeExporter.export_pdf_to_excel = _improved_export_pdf_to_excel
except NameError:
    pass



# --- word korean font fix override v1 ---
def _set_docx_run_font(run, font_name: str = "Malgun Gothic", font_size_pt: int | None = None) -> None:
    """
    Apply Korean/English font to a python-docx run.
    East Asian font must be set separately, otherwise Korean can render incorrectly.
    """
    try:
        from docx.shared import Pt
        from docx.oxml.ns import qn

        run.font.name = font_name

        if font_size_pt:
            run.font.size = Pt(font_size_pt)

        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts

        if r_fonts is None:
            from docx.oxml import OxmlElement
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)

        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)
        r_fonts.set(qn("w:eastAsia"), font_name)
        r_fonts.set(qn("w:cs"), font_name)

    except Exception:
        pass


def _set_docx_default_korean_font(docx, font_name: str = "Malgun Gothic") -> None:
    """
    Set document default styles to Korean-friendly font.
    """
    try:
        from docx.shared import Pt
        from docx.oxml.ns import qn

        styles = docx.styles

        for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
            try:
                style = styles[style_name]
                style.font.name = font_name

                if style_name == "Normal":
                    style.font.size = Pt(10)
                elif style_name == "Heading 1":
                    style.font.size = Pt(16)
                elif style_name == "Heading 2":
                    style.font.size = Pt(13)

                r_pr = style._element.get_or_add_rPr()
                r_fonts = r_pr.rFonts

                if r_fonts is None:
                    from docx.oxml import OxmlElement
                    r_fonts = OxmlElement("w:rFonts")
                    r_pr.append(r_fonts)

                r_fonts.set(qn("w:ascii"), font_name)
                r_fonts.set(qn("w:hAnsi"), font_name)
                r_fonts.set(qn("w:eastAsia"), font_name)
                r_fonts.set(qn("w:cs"), font_name)

            except Exception:
                continue

    except Exception:
        pass


def _add_docx_paragraph_with_korean_font(docx, text: str, font_name: str = "Malgun Gothic") -> None:
    paragraph = docx.add_paragraph()
    run = paragraph.add_run(text)
    _set_docx_run_font(run, font_name=font_name, font_size_pt=10)


def _word_korean_font_export_pdf_to_word(
    self,
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Word export with Korean font support.
    """
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    import tempfile

    options = options or OfficeExportOptions()
    word_mode = getattr(options, "word_mode", "text")

    source_path = Path(source_pdf_path)
    output_path = Path(output_docx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Word export.")

        docx = Document()
        _set_docx_default_korean_font(docx, "Malgun Gothic")

        title = docx.add_heading(source_path.name, level=1)
        for run in title.runs:
            _set_docx_run_font(run, "Malgun Gothic", 16)

        with tempfile.TemporaryDirectory(prefix="pdf_docx_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)

                heading = docx.add_heading(f"Page {page_index + 1}", level=2)
                for run in heading.runs:
                    _set_docx_run_font(run, "Malgun Gothic", 13)

                if word_mode == "image_text":
                    try:
                        zoom = max(float(getattr(options, "dpi", 200)), 72.0) / 72.0
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(zoom, zoom),
                            alpha=False,
                            colorspace=fitz.csRGB,
                        )
                        image_path = temp_dir_path / f"page_{page_index + 1:04d}.png"
                        pix.save(str(image_path))
                        docx.add_picture(str(image_path), width=Inches(6.5))
                    except Exception:
                        pass

                    docx.add_paragraph("")

                text = page.get_text("text").strip()
                text = _fix_broken_pdf_symbols(text)

                if text:
                    for paragraph_text in re.split(r"\n\s*\n", text):
                        paragraph_text = paragraph_text.strip()

                        if not paragraph_text:
                            continue

                        paragraph_text = _aggressive_korean_spacing(
                            _final_korean_spacing_cleanup(
                                _apply_korean_spacing_heuristics(paragraph_text)
                            )
                        )
                        paragraph_text = _remove_leaked_brand_tokens(paragraph_text)
                        paragraph_text = _fix_broken_pdf_symbols(paragraph_text)

                        _add_docx_paragraph_with_korean_font(
                            docx,
                            paragraph_text,
                            "Malgun Gothic",
                        )
                else:
                    _add_docx_paragraph_with_korean_font(
                        docx,
                        "[No extractable text on this page]",
                        "Malgun Gothic",
                    )

                if options.include_page_breaks and order != len(page_indexes) - 1:
                    docx.add_page_break()

        docx.save(str(output_path))

    finally:
        pdf.close()


try:
    OfficeExporter.export_pdf_to_word = _word_korean_font_export_pdf_to_word
except NameError:
    pass



# --- word layout image mode override v1 ---
def _set_word_section_to_pdf_page(section, page_rect) -> None:
    """
    Set Word page size close to PDF page size.
    """
    try:
        from docx.shared import Pt

        section.page_width = Pt(float(page_rect.width))
        section.page_height = Pt(float(page_rect.height))

        margin = Pt(0)
        section.top_margin = margin
        section.bottom_margin = margin
        section.left_margin = margin
        section.right_margin = margin
    except Exception:
        pass


def _word_layout_image_export_pdf_to_word(
    self,
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Word export with layout-preserving image mode.

    Modes:
    - text: text only
    - image_text: page image + extracted text
    - layout_image: page image only, sized to Word page
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from exc

    import tempfile

    options = options or OfficeExportOptions()
    word_mode = getattr(options, "word_mode", "text")

    source_path = Path(source_pdf_path)
    output_path = Path(output_docx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Word export.")

        docx = Document()
        _set_docx_default_korean_font(docx, "Malgun Gothic")

        with tempfile.TemporaryDirectory(prefix="pdf_docx_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            first_page_done = False

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)

                if word_mode == "layout_image":
                    if not first_page_done:
                        _set_word_section_to_pdf_page(docx.sections[0], page.rect)
                        first_page_done = True
                    elif order > 0:
                        docx.add_page_break()

                    zoom = max(float(getattr(options, "dpi", 200)), 72.0) / 72.0
                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(zoom, zoom),
                        alpha=False,
                        colorspace=fitz.csRGB,
                    )

                    image_path = temp_dir_path / f"page_{page_index + 1:04d}.png"
                    pix.save(str(image_path))

                    paragraph = docx.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = paragraph.add_run()
                    run.add_picture(str(image_path), width=Pt(float(page.rect.width)))

                    continue

                # Existing text/image_text behavior
                heading = docx.add_heading(f"Page {page_index + 1}", level=2)
                for run in heading.runs:
                    _set_docx_run_font(run, "Malgun Gothic", 13)

                if word_mode == "image_text":
                    try:
                        zoom = max(float(getattr(options, "dpi", 200)), 72.0) / 72.0
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(zoom, zoom),
                            alpha=False,
                            colorspace=fitz.csRGB,
                        )
                        image_path = temp_dir_path / f"page_{page_index + 1:04d}.png"
                        pix.save(str(image_path))
                        docx.add_picture(str(image_path), width=Inches(6.5))
                    except Exception:
                        pass

                    docx.add_paragraph("")

                text = page.get_text("text").strip()
                text = _fix_broken_pdf_symbols(text)

                if text:
                    for paragraph_text in re.split(r"\n\s*\n", text):
                        paragraph_text = paragraph_text.strip()

                        if not paragraph_text:
                            continue

                        paragraph_text = _aggressive_korean_spacing(
                            _final_korean_spacing_cleanup(
                                _apply_korean_spacing_heuristics(paragraph_text)
                            )
                        )
                        paragraph_text = _remove_leaked_brand_tokens(paragraph_text)
                        paragraph_text = _fix_broken_pdf_symbols(paragraph_text)

                        _add_docx_paragraph_with_korean_font(
                            docx,
                            paragraph_text,
                            "Malgun Gothic",
                        )
                else:
                    _add_docx_paragraph_with_korean_font(
                        docx,
                        "[No extractable text on this page]",
                        "Malgun Gothic",
                    )

                if options.include_page_breaks and order != len(page_indexes) - 1:
                    docx.add_page_break()

        docx.save(str(output_path))

    finally:
        pdf.close()


try:
    OfficeExporter.export_pdf_to_word = _word_layout_image_export_pdf_to_word
except NameError:
    pass



# --- word com editable layout override v1 ---
def _export_pdf_to_word_com_editable_layout(
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Export PDF to editable Word layout using Microsoft Word COM.

    Requirements:
    - Windows
    - Microsoft Word installed
    - pywin32 installed

    Strategy:
    - Set Word page size to PDF page size.
    - Add PDF text as absolutely positioned Word text boxes.
    - Add embedded PDF image blocks as absolutely positioned Word pictures.
    """
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("pywin32 is not installed. Run: pip install pywin32") from exc

    import tempfile

    options = options or OfficeExportOptions()
    source_path = Path(source_pdf_path)
    output_path = Path(output_docx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))
    word = None
    doc = None

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Word export.")

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Add()

        # Word constants
        wdStory = 6
        wdPageBreak = 7
        msoTextOrientationHorizontal = 1
        msoFalse = 0
        msoTrue = -1

        with tempfile.TemporaryDirectory(prefix="pdf_word_com_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                page_rect = page.rect

                # New page except first.
                if order > 0:
                    word.Selection.EndKey(Unit=wdStory)
                    word.Selection.InsertBreak(Type=wdPageBreak)

                section = doc.Sections(doc.Sections.Count)
                section.PageSetup.PageWidth = float(page_rect.width)
                section.PageSetup.PageHeight = float(page_rect.height)
                section.PageSetup.TopMargin = 0
                section.PageSetup.BottomMargin = 0
                section.PageSetup.LeftMargin = 0
                section.PageSetup.RightMargin = 0

                # Add embedded PDF images first.
                try:
                    data = page.get_text("dict")
                    for block_index, block in enumerate(data.get("blocks", [])):
                        if block.get("type") != 1:
                            continue

                        bbox = block.get("bbox")
                        image_bytes = block.get("image")

                        if not bbox or len(bbox) != 4 or not image_bytes:
                            continue

                        x0, y0, x1, y1 = [float(v) for v in bbox]
                        img_path = temp_dir_path / f"p{page_index + 1}_img{block_index}.png"
                        img_path.write_bytes(image_bytes)

                        shape = doc.Shapes.AddPicture(
                            FileName=str(img_path),
                            LinkToFile=msoFalse,
                            SaveWithDocument=msoTrue,
                            Left=x0,
                            Top=y0,
                            Width=max(x1 - x0, 1),
                            Height=max(y1 - y0, 1),
                        )

                        shape.WrapFormat.Type = 3  # wdWrapNone-ish / floating
                except Exception:
                    pass

                # Add text as positioned text boxes using word-based lines.
                lines = []
                try:
                    lines = _group_pdf_words_into_lines(page)
                except Exception:
                    lines = []

                if not lines:
                    # Fallback: make simple full-page text box.
                    text = page.get_text("text").strip()
                    if text:
                        text = _fix_broken_pdf_symbols(text)
                        text = _aggressive_korean_spacing(
                            _final_korean_spacing_cleanup(
                                _apply_korean_spacing_heuristics(text)
                            )
                        )
                        text = _remove_leaked_brand_tokens(text)

                        shape = doc.Shapes.AddTextbox(
                            msoTextOrientationHorizontal,
                            20,
                            20,
                            max(float(page_rect.width) - 40, 50),
                            max(float(page_rect.height) - 40, 50),
                        )
                        shape.TextFrame.TextRange.Text = text
                        shape.TextFrame.TextRange.Font.Name = "Malgun Gothic"
                        shape.TextFrame.TextRange.Font.Size = 10
                        shape.Line.Visible = False
                        shape.Fill.Visible = False
                    continue

                for line in lines:
                    try:
                        text = line.get("text", "")
                        if not text:
                            continue

                        x0, y0, x1, y1 = line.get("bbox", (0, 0, 100, 20))
                        avg_height = float(line.get("avg_height", max(y1 - y0, 10.0)))

                        # Slightly enlarge box to reduce clipping.
                        left = float(x0)
                        top = float(y0)
                        width = max(float(x1 - x0) + 8, 12)
                        height = max(float(y1 - y0) + 6, 10)

                        shape = doc.Shapes.AddTextbox(
                            msoTextOrientationHorizontal,
                            left,
                            top,
                            width,
                            height,
                        )

                        shape.TextFrame.MarginLeft = 0
                        shape.TextFrame.MarginRight = 0
                        shape.TextFrame.MarginTop = 0
                        shape.TextFrame.MarginBottom = 0
                        shape.TextFrame.WordWrap = False

                        tr = shape.TextFrame.TextRange
                        tr.Text = text
                        tr.Font.Name = "Malgun Gothic"
                        tr.Font.Size = max(avg_height * 0.72, 5)

                        # Transparent box.
                        shape.Line.Visible = False
                        shape.Fill.Visible = False

                    except Exception:
                        continue

        doc.SaveAs(str(output_path.resolve()))

    finally:
        try:
            if doc is not None:
                doc.Close(SaveChanges=False)
        except Exception:
            pass

        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass

        pdf.close()


def _word_com_layout_export_pdf_to_word(
    self,
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    options = options or OfficeExportOptions()

    if getattr(options, "word_mode", "text") == "layout_editable_com":
        return _export_pdf_to_word_com_editable_layout(
            source_pdf_path,
            output_docx_path,
            options,
        )

    # Fallback to previous active implementation.
    return _previous_word_export_for_com_layout(
        self,
        source_pdf_path,
        output_docx_path,
        options,
    )


try:
    if not hasattr(OfficeExporter, "_previous_word_export_for_com_layout"):
        OfficeExporter._previous_word_export_for_com_layout = OfficeExporter.export_pdf_to_word

    _previous_word_export_for_com_layout = OfficeExporter._previous_word_export_for_com_layout
    OfficeExporter.export_pdf_to_word = _word_com_layout_export_pdf_to_word
except NameError:
    pass



# --- word com drawing shapes override v1 ---
def _word_color_int_from_pdf_color(color_value, default=(0, 0, 0)) -> int:
    """
    Convert PDF RGB tuple 0~1 to Word COM BGR integer.
    Word uses BGR: R + G*256 + B*65536.
    """
    try:
        if color_value is None:
            r, g, b = default
        elif isinstance(color_value, (list, tuple)) and len(color_value) >= 3:
            r = int(max(0.0, min(1.0, float(color_value[0]))) * 255)
            g = int(max(0.0, min(1.0, float(color_value[1]))) * 255)
            b = int(max(0.0, min(1.0, float(color_value[2]))) * 255)
        else:
            r, g, b = default
    except Exception:
        r, g, b = default

    return int(r + (g * 256) + (b * 65536))


def _add_pdf_drawings_to_word_doc(
    doc,
    page: fitz.Page,
    max_drawing_objects: int = 800,
) -> None:
    """
    Add simple PDF drawings to Word document using COM Shapes.

    Supported:
    - straight lines
    - rectangles
    - thin rectangles as lines
    """
    if max_drawing_objects <= 0:
        return

    try:
        drawings = page.get_drawings()
    except Exception:
        return

    if not drawings:
        return

    # Word constants
    msoShapeRectangle = 1
    msoLine = 9
    msoFalse = 0
    msoTrue = -1

    page_area = max(float(page.rect.width * page.rect.height), 1.0)

    def add_line(x0, y0, x1, y1, color, width):
        try:
            shape = doc.Shapes.AddLine(float(x0), float(y0), float(x1), float(y1))
            shape.Line.Visible = msoTrue
            shape.Line.ForeColor.RGB = _word_color_int_from_pdf_color(color)
            shape.Line.Weight = max(float(width or 0.5) * 1.15, 0.35)
            return shape
        except Exception:
            return None

    def add_rect(x0, y0, x1, y1, stroke_color, fill_color, width):
        try:
            left = min(float(x0), float(x1))
            top = min(float(y0), float(y1))
            w = abs(float(x1) - float(x0))
            h = abs(float(y1) - float(y0))

            if w <= 1 and h <= 1:
                return None

            # Thin rects usually represent lines.
            if w <= 2 and h > 2:
                return add_line(left, top, left, top + h, stroke_color or fill_color, width)
            if h <= 2 and w > 2:
                return add_line(left, top, left + w, top, stroke_color or fill_color, width)

            shape = doc.Shapes.AddShape(
                msoShapeRectangle,
                left,
                top,
                max(w, 1),
                max(h, 1),
            )

            if fill_color is not None:
                shape.Fill.Visible = msoTrue
                shape.Fill.ForeColor.RGB = _word_color_int_from_pdf_color(fill_color, default=(255, 255, 255))
            else:
                shape.Fill.Visible = msoFalse

            if stroke_color is not None:
                shape.Line.Visible = msoTrue
                shape.Line.ForeColor.RGB = _word_color_int_from_pdf_color(stroke_color)
                shape.Line.Weight = max(float(width or 0.5) * 1.15, 0.35)
            else:
                shape.Line.Visible = msoFalse

            return shape

        except Exception:
            return None

    converted = 0

    for drawing in drawings:
        if converted >= max_drawing_objects:
            break

        try:
            stroke_color = drawing.get("color")
            fill_color = drawing.get("fill")
            line_width = float(drawing.get("width", 0.5) or 0.5)
            rect = drawing.get("rect")
            items = drawing.get("items", [])

            if rect is not None:
                area = max(float(rect.width * rect.height), 0.0)
                # Avoid huge background blocks.
                if area > page_area * 0.6:
                    continue

            # Skip overly complex paths for Word COM stability.
            if len(items) > 50:
                continue

            handled = False

            for item in items:
                if not item:
                    continue

                op = item[0]

                if op == "re" and len(item) >= 2:
                    r = item[1]
                    shape = add_rect(
                        r.x0,
                        r.y0,
                        r.x1,
                        r.y1,
                        stroke_color,
                        fill_color,
                        line_width,
                    )
                    if shape is not None:
                        converted += 1
                        handled = True

                elif op == "l" and len(item) >= 3:
                    p1 = item[1]
                    p2 = item[2]

                    if abs(float(p1.x) - float(p2.x)) < 0.2 and abs(float(p1.y) - float(p2.y)) < 0.2:
                        continue

                    shape = add_line(
                        p1.x,
                        p1.y,
                        p2.x,
                        p2.y,
                        stroke_color,
                        line_width,
                    )
                    if shape is not None:
                        converted += 1
                        handled = True

                if converted >= max_drawing_objects:
                    break

            # Fallback filled rectangle from drawing bbox.
            if not handled and rect is not None and fill_color is not None:
                shape = add_rect(
                    rect.x0,
                    rect.y0,
                    rect.x1,
                    rect.y1,
                    stroke_color,
                    fill_color,
                    line_width,
                )
                if shape is not None:
                    converted += 1

        except Exception:
            continue


_previous_word_com_layout_export = _export_pdf_to_word_com_editable_layout


def _export_pdf_to_word_com_editable_layout(
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Word COM editable layout export with drawing support.

    This replaces the earlier COM layout function.
    """
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("pywin32 is not installed. Run: pip install pywin32") from exc

    import tempfile

    options = options or OfficeExportOptions()
    source_path = Path(source_pdf_path)
    output_path = Path(output_docx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))
    word = None
    doc = None

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Word export.")

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Add()

        wdStory = 6
        wdPageBreak = 7
        msoTextOrientationHorizontal = 1
        msoFalse = 0
        msoTrue = -1

        with tempfile.TemporaryDirectory(prefix="pdf_word_com_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                page_rect = page.rect

                if order > 0:
                    word.Selection.EndKey(Unit=wdStory)
                    word.Selection.InsertBreak(Type=wdPageBreak)

                section = doc.Sections(doc.Sections.Count)
                section.PageSetup.PageWidth = float(page_rect.width)
                section.PageSetup.PageHeight = float(page_rect.height)
                section.PageSetup.TopMargin = 0
                section.PageSetup.BottomMargin = 0
                section.PageSetup.LeftMargin = 0
                section.PageSetup.RightMargin = 0

                # Drawings first.
                try:
                    _add_pdf_drawings_to_word_doc(
                        doc,
                        page,
                        max_drawing_objects=getattr(options, "max_drawing_objects", 1200),
                        temp_dir_path=temp_dir_path,
                    )
                except Exception:
                    pass

                # Images second.
                try:
                    data = page.get_text("dict")
                    for block_index, block in enumerate(data.get("blocks", [])):
                        if block.get("type") != 1:
                            continue

                        bbox = block.get("bbox")
                        image_bytes = block.get("image")

                        if not bbox or len(bbox) != 4 or not image_bytes:
                            continue

                        x0, y0, x1, y1 = [float(v) for v in bbox]
                        img_path = temp_dir_path / f"p{page_index + 1}_img{block_index}.png"
                        img_path.write_bytes(image_bytes)

                        shape = doc.Shapes.AddPicture(
                            FileName=str(img_path),
                            LinkToFile=msoFalse,
                            SaveWithDocument=msoTrue,
                            Left=x0,
                            Top=y0,
                            Width=max(x1 - x0, 1),
                            Height=max(y1 - y0, 1),
                        )
                        shape.WrapFormat.Type = 3
                except Exception:
                    pass

                # Text last, so it stays on top.
                try:
                    lines = _group_pdf_words_into_lines(page)
                except Exception:
                    lines = []

                if not lines:
                    text = page.get_text("text").strip()
                    if text:
                        text = _fix_broken_pdf_symbols(text)
                        text = _aggressive_korean_spacing(
                            _final_korean_spacing_cleanup(
                                _apply_korean_spacing_heuristics(text)
                            )
                        )
                        text = _remove_leaked_brand_tokens(text)

                        shape = doc.Shapes.AddTextbox(
                            msoTextOrientationHorizontal,
                            20,
                            20,
                            max(float(page_rect.width) - 40, 50),
                            max(float(page_rect.height) - 40, 50),
                        )
                        shape.TextFrame.TextRange.Text = text
                        shape.TextFrame.TextRange.Font.Name = "Malgun Gothic"
                        shape.TextFrame.TextRange.Font.Size = 10
                        shape.Line.Visible = False
                        shape.Fill.Visible = False
                    continue

                for line in lines:
                    try:
                        text = line.get("text", "")
                        if not text:
                            continue

                        x0, y0, x1, y1 = line.get("bbox", (0, 0, 100, 20))
                        avg_height = float(line.get("avg_height", max(y1 - y0, 10.0)))

                        shape = doc.Shapes.AddTextbox(
                            msoTextOrientationHorizontal,
                            float(x0),
                            float(y0),
                            max(float(x1 - x0) + 8, 12),
                            max(float(y1 - y0) + 6, 10),
                        )

                        shape.TextFrame.MarginLeft = 0
                        shape.TextFrame.MarginRight = 0
                        shape.TextFrame.MarginTop = 0
                        shape.TextFrame.MarginBottom = 0
                        shape.TextFrame.WordWrap = False

                        tr = shape.TextFrame.TextRange
                        tr.Text = text
                        tr.Font.Name = "Malgun Gothic"
                        tr.Font.Size = max(avg_height * 0.72, 5)

                        shape.Line.Visible = False
                        shape.Fill.Visible = False

                    except Exception:
                        continue

        doc.SaveAs(str(output_path.resolve()))

    finally:
        try:
            if doc is not None:
                doc.Close(SaveChanges=False)
        except Exception:
            pass

        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass

        pdf.close()



# --- word com complex drawing fallback override v1 ---
def _boost_rect_for_word_fallback(rect: fitz.Rect, amount: float, page_rect: fitz.Rect) -> fitz.Rect:
    return fitz.Rect(
        max(page_rect.x0, rect.x0 - amount),
        max(page_rect.y0, rect.y0 - amount),
        min(page_rect.x1, rect.x1 + amount),
        min(page_rect.y1, rect.y1 + amount),
    )


def _add_pdf_clip_image_to_word_doc(
    doc,
    page: fitz.Page,
    clip: fitz.Rect,
    temp_dir_path: Path,
    file_prefix: str,
    dpi: int = 220,
) -> bool:
    """
    Render a PDF clip as image and add it to Word at the same coordinates.
    """
    try:
        if clip.width <= 0 or clip.height <= 0:
            return False

        zoom = max(float(dpi), 96.0) / 72.0
        pix = page.get_pixmap(
            matrix=fitz.Matrix(zoom, zoom),
            clip=clip,
            alpha=True,
            colorspace=fitz.csRGB,
        )

        image_path = temp_dir_path / f"{file_prefix}.png"
        pix.save(str(image_path))

        # Word constants
        msoFalse = 0
        msoTrue = -1

        shape = doc.Shapes.AddPicture(
            FileName=str(image_path),
            LinkToFile=msoFalse,
            SaveWithDocument=msoTrue,
            Left=float(clip.x0),
            Top=float(clip.y0),
            Width=max(float(clip.width), 1.0),
            Height=max(float(clip.height), 1.0),
        )

        try:
            shape.WrapFormat.Type = 3
        except Exception:
            pass

        return True

    except Exception:
        return False


def _add_pdf_drawings_to_word_doc(
    doc,
    page: fitz.Page,
    max_drawing_objects: int = 1200,
    temp_dir_path: Path | None = None,
) -> None:
    """
    Word drawing engine v2.

    Supported:
    - simple lines as Word lines
    - rectangles as Word rectangles
    - complex small vector drawings as clipped images
    """
    if max_drawing_objects <= 0:
        return

    import tempfile

    owns_temp_dir = False
    temp_ctx = None

    if temp_dir_path is None:
        temp_ctx = tempfile.TemporaryDirectory(prefix="word_drawing_fallback_")
        temp_dir_path = Path(temp_ctx.name)
        owns_temp_dir = True

    try:
        drawings = page.get_drawings()
    except Exception:
        if temp_ctx is not None:
            temp_ctx.cleanup()
        return

    if not drawings:
        if temp_ctx is not None:
            temp_ctx.cleanup()
        return

    # Word constants
    msoShapeRectangle = 1
    msoFalse = 0
    msoTrue = -1

    page_rect = page.rect
    page_area = max(float(page.rect.width * page.rect.height), 1.0)

    def add_line(x0, y0, x1, y1, color, width):
        try:
            shape = doc.Shapes.AddLine(float(x0), float(y0), float(x1), float(y1))
            shape.Line.Visible = msoTrue
            shape.Line.ForeColor.RGB = _word_color_int_from_pdf_color(color)
            shape.Line.Weight = max(float(width or 0.5) * 1.15, 0.35)
            return shape
        except Exception:
            return None

    def add_rect(x0, y0, x1, y1, stroke_color, fill_color, width):
        try:
            left = min(float(x0), float(x1))
            top = min(float(y0), float(y1))
            w = abs(float(x1) - float(x0))
            h = abs(float(y1) - float(y0))

            if w <= 1 and h <= 1:
                return None

            if w <= 2 and h > 2:
                return add_line(left, top, left, top + h, stroke_color or fill_color, width)

            if h <= 2 and w > 2:
                return add_line(left, top, left + w, top, stroke_color or fill_color, width)

            shape = doc.Shapes.AddShape(
                msoShapeRectangle,
                left,
                top,
                max(w, 1),
                max(h, 1),
            )

            if fill_color is not None:
                shape.Fill.Visible = msoTrue
                shape.Fill.ForeColor.RGB = _word_color_int_from_pdf_color(
                    fill_color,
                    default=(255, 255, 255),
                )
            else:
                shape.Fill.Visible = msoFalse

            if stroke_color is not None:
                shape.Line.Visible = msoTrue
                shape.Line.ForeColor.RGB = _word_color_int_from_pdf_color(stroke_color)
                shape.Line.Weight = max(float(width or 0.5) * 1.15, 0.35)
            else:
                shape.Line.Visible = msoFalse

            return shape

        except Exception:
            return None

    converted = 0
    fallback_images = 0
    max_fallback_images = 180

    try:
        for drawing_index, drawing in enumerate(drawings):
            if converted >= max_drawing_objects:
                break

            try:
                stroke_color = drawing.get("color")
                fill_color = drawing.get("fill")
                line_width = float(drawing.get("width", 0.5) or 0.5)
                rect = drawing.get("rect")
                items = drawing.get("items", [])

                if rect is None:
                    continue

                clip = fitz.Rect(rect)
                clip_area = max(float(clip.width * clip.height), 0.0)

                # Avoid huge page/background fallback.
                if clip_area <= 0 or clip_area > page_area * 0.35:
                    continue

                if clip.width < 0.8 or clip.height < 0.8:
                    continue

                handled = False
                complex_count = 0

                # Keep very complex paths as image fallback, not individual Word shapes.
                if len(items) <= 50:
                    for item in items:
                        if not item:
                            continue

                        op = item[0]

                        if op == "re" and len(item) >= 2:
                            r = item[1]
                            shape = add_rect(
                                r.x0,
                                r.y0,
                                r.x1,
                                r.y1,
                                stroke_color,
                                fill_color,
                                line_width,
                            )

                            if shape is not None:
                                converted += 1
                                handled = True

                        elif op == "l" and len(item) >= 3:
                            p1 = item[1]
                            p2 = item[2]

                            if abs(float(p1.x) - float(p2.x)) < 0.2 and abs(float(p1.y) - float(p2.y)) < 0.2:
                                continue

                            shape = add_line(
                                p1.x,
                                p1.y,
                                p2.x,
                                p2.y,
                                stroke_color,
                                line_width,
                            )

                            if shape is not None:
                                converted += 1
                                handled = True

                        else:
                            complex_count += 1

                        if converted >= max_drawing_objects:
                            break
                else:
                    complex_count = len(items)

                if handled:
                    continue

                # Simple filled rectangle fallback.
                if fill_color is not None and len(items) <= 4:
                    shape = add_rect(
                        clip.x0,
                        clip.y0,
                        clip.x1,
                        clip.y1,
                        stroke_color,
                        fill_color,
                        line_width,
                    )

                    if shape is not None:
                        converted += 1
                        continue

                # Complex small/medium vector fallback as image.
                if fallback_images < max_fallback_images:
                    is_small_icon = clip_area < page_area * 0.04
                    is_medium_graphic = clip_area < page_area * 0.12 and len(items) >= 3

                    if complex_count > 0 and (is_small_icon or is_medium_graphic):
                        boosted = _boost_rect_for_word_fallback(clip, 1.0, page_rect)

                        ok = _add_pdf_clip_image_to_word_doc(
                            doc=doc,
                            page=page,
                            clip=boosted,
                            temp_dir_path=temp_dir_path,
                            file_prefix=f"word_p{page.number + 1}_draw_{drawing_index}",
                            dpi=220,
                        )

                        if ok:
                            fallback_images += 1
                            converted += 1

            except Exception:
                continue

    finally:
        if temp_ctx is not None:
            try:
                temp_ctx.cleanup()
            except Exception:
                pass



# --- word auto mode override v1 ---
def _export_pdf_to_word_auto_layout(
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    """
    Word auto mode.

    Per page:
    - complex technical/vector pages -> image page
    - editable-friendly pages -> Word COM editable layout
    """
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("pywin32 is not installed. Run: pip install pywin32") from exc

    import tempfile

    options = options or OfficeExportOptions()
    source_path = Path(source_pdf_path)
    output_path = Path(output_docx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))
    word = None
    doc = None

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Word export.")

        options.word_auto_report = []

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Add()

        wdStory = 6
        wdPageBreak = 7
        msoTextOrientationHorizontal = 1
        msoFalse = 0
        msoTrue = -1

        with tempfile.TemporaryDirectory(prefix="pdf_word_auto_") as temp_dir:
            temp_dir_path = Path(temp_dir)

            for order, page_index in enumerate(page_indexes):
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                page_rect = page.rect

                complexity = _analyze_pdf_page_complexity_for_ppt(page)
                effective_mode = "image" if complexity.get("complex") else "editable"

                try:
                    options.word_auto_report.append(
                        {
                            "page": page_index + 1,
                            "mode": effective_mode,
                            "reason": complexity.get("reason", ""),
                            "drawings": complexity.get("drawings", 0),
                            "drawing_items": complexity.get("drawing_items", 0),
                            "words": complexity.get("words", 0),
                        }
                    )

                    print(
                        f"[WORD AUTO] Page {page_index + 1}: {effective_mode} - "
                        f"{complexity.get('reason', '')}"
                    )
                except Exception:
                    pass

                if order > 0:
                    word.Selection.EndKey(Unit=wdStory)
                    word.Selection.InsertBreak(Type=wdPageBreak)

                section = doc.Sections(doc.Sections.Count)
                section.PageSetup.PageWidth = float(page_rect.width)
                section.PageSetup.PageHeight = float(page_rect.height)
                section.PageSetup.TopMargin = 0
                section.PageSetup.BottomMargin = 0
                section.PageSetup.LeftMargin = 0
                section.PageSetup.RightMargin = 0

                if effective_mode == "image":
                    zoom = max(float(getattr(options, "dpi", 200)), 72.0) / 72.0
                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(zoom, zoom),
                        alpha=False,
                        colorspace=fitz.csRGB,
                    )

                    image_path = temp_dir_path / f"word_auto_page_{page_index + 1:04d}.png"
                    pix.save(str(image_path))

                    shape = doc.Shapes.AddPicture(
                        FileName=str(image_path),
                        LinkToFile=msoFalse,
                        SaveWithDocument=msoTrue,
                        Left=0,
                        Top=0,
                        Width=float(page_rect.width),
                        Height=float(page_rect.height),
                    )

                    try:
                        shape.WrapFormat.Type = 3
                    except Exception:
                        pass

                    continue

                # Editable mode: drawings/images/text.
                try:
                    _add_pdf_drawings_to_word_doc(
                        doc,
                        page,
                        max_drawing_objects=getattr(options, "max_drawing_objects", 1200),
                        temp_dir_path=temp_dir_path,
                    )
                except Exception:
                    pass

                try:
                    data = page.get_text("dict")
                    for block_index, block in enumerate(data.get("blocks", [])):
                        if block.get("type") != 1:
                            continue

                        bbox = block.get("bbox")
                        image_bytes = block.get("image")

                        if not bbox or len(bbox) != 4 or not image_bytes:
                            continue

                        x0, y0, x1, y1 = [float(v) for v in bbox]
                        img_path = temp_dir_path / f"word_auto_p{page_index + 1}_img{block_index}.png"
                        img_path.write_bytes(image_bytes)

                        shape = doc.Shapes.AddPicture(
                            FileName=str(img_path),
                            LinkToFile=msoFalse,
                            SaveWithDocument=msoTrue,
                            Left=x0,
                            Top=y0,
                            Width=max(x1 - x0, 1),
                            Height=max(y1 - y0, 1),
                        )

                        try:
                            shape.WrapFormat.Type = 3
                        except Exception:
                            pass

                except Exception:
                    pass

                try:
                    lines = _group_pdf_words_into_lines(page)
                except Exception:
                    lines = []

                if not lines:
                    text = page.get_text("text").strip()

                    if text:
                        text = _fix_broken_pdf_symbols(text)
                        text = _aggressive_korean_spacing(
                            _final_korean_spacing_cleanup(
                                _apply_korean_spacing_heuristics(text)
                            )
                        )
                        text = _remove_leaked_brand_tokens(text)

                        shape = doc.Shapes.AddTextbox(
                            msoTextOrientationHorizontal,
                            20,
                            20,
                            max(float(page_rect.width) - 40, 50),
                            max(float(page_rect.height) - 40, 50),
                        )
                        shape.TextFrame.TextRange.Text = text
                        shape.TextFrame.TextRange.Font.Name = "Malgun Gothic"
                        shape.TextFrame.TextRange.Font.Size = 10
                        shape.Line.Visible = False
                        shape.Fill.Visible = False

                    continue

                for line in lines:
                    try:
                        text = line.get("text", "")
                        if not text:
                            continue

                        x0, y0, x1, y1 = line.get("bbox", (0, 0, 100, 20))
                        avg_height = float(line.get("avg_height", max(y1 - y0, 10.0)))

                        shape = doc.Shapes.AddTextbox(
                            msoTextOrientationHorizontal,
                            float(x0),
                            float(y0),
                            max(float(x1 - x0) + 8, 12),
                            max(float(y1 - y0) + 6, 10),
                        )

                        shape.TextFrame.MarginLeft = 0
                        shape.TextFrame.MarginRight = 0
                        shape.TextFrame.MarginTop = 0
                        shape.TextFrame.MarginBottom = 0
                        shape.TextFrame.WordWrap = False

                        tr = shape.TextFrame.TextRange
                        tr.Text = text
                        tr.Font.Name = "Malgun Gothic"
                        tr.Font.Size = max(avg_height * 0.72, 5)

                        shape.Line.Visible = False
                        shape.Fill.Visible = False

                    except Exception:
                        continue

        doc.SaveAs(str(output_path.resolve()))

    finally:
        try:
            if doc is not None:
                doc.Close(SaveChanges=False)
        except Exception:
            pass

        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass

        pdf.close()


def _word_auto_mode_export_pdf_to_word(
    self,
    source_pdf_path: str,
    output_docx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    options = options or OfficeExportOptions()

    if getattr(options, "word_mode", "text") == "auto_layout":
        return _export_pdf_to_word_auto_layout(
            source_pdf_path,
            output_docx_path,
            options,
        )

    return _previous_word_export_for_auto_mode(
        self,
        source_pdf_path,
        output_docx_path,
        options,
    )


try:
    if not hasattr(OfficeExporter, "_previous_word_export_for_auto_mode"):
        OfficeExporter._previous_word_export_for_auto_mode = OfficeExporter.export_pdf_to_word

    _previous_word_export_for_auto_mode = OfficeExporter._previous_word_export_for_auto_mode
    OfficeExporter.export_pdf_to_word = _word_auto_mode_export_pdf_to_word
except NameError:
    pass



# --- excel table quality override v1 ---
def _autosize_excel_columns(ws, max_width: int = 60) -> None:
    try:
        from openpyxl.utils import get_column_letter

        for column_cells in ws.columns:
            max_len = 0
            col_idx = column_cells[0].column

            for cell in column_cells:
                value = cell.value
                if value is None:
                    continue
                max_len = max(max_len, len(str(value)))

            width = min(max(max_len + 2, 10), max_width)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    except Exception:
        pass


def _style_excel_sheet(ws) -> None:
    try:
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        default_font = Font(name="Malgun Gothic", size=10)
        header_font = Font(name="Malgun Gothic", size=10, bold=True)
        header_fill = PatternFill("solid", fgColor="EDEDED")
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row in ws.iter_rows():
            for cell in row:
                cell.font = default_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

        # First row as header/title.
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        ws.freeze_panes = "A2"

    except Exception:
        pass


def _improved_group_words_to_table_rows_v2(page: fitz.Page) -> list[list[str]]:
    """
    Better coordinate-based table reconstruction.

    Strategy:
    - group words into visual rows by y-center
    - split cells by horizontal gaps
    - keep Korean spacing cleanup
    """
    try:
        words = page.get_text("words")
    except Exception:
        return []

    items = []

    for w in words:
        if len(w) < 8:
            continue

        x0, y0, x1, y1, text, block_no, line_no, word_no = w[:8]
        text = _fix_broken_pdf_symbols(str(text or "")).strip()

        if not text:
            continue

        x0 = float(x0)
        y0 = float(y0)
        x1 = float(x1)
        y1 = float(y1)
        h = max(y1 - y0, 1.0)
        y_center = (y0 + y1) / 2.0

        items.append(
            {
                "x0": x0,
                "y0": y0,
                "x1": x1,
                "y1": y1,
                "yc": y_center,
                "h": h,
                "text": text,
            }
        )

    if not items:
        return []

    avg_h = sum(it["h"] for it in items) / len(items)
    row_threshold = max(avg_h * 0.65, 3.5)

    items.sort(key=lambda it: (it["yc"], it["x0"]))

    rows: list[list[dict]] = []

    for item in items:
        placed = False

        for row in rows:
            row_y = sum(it["yc"] for it in row) / len(row)

            if abs(item["yc"] - row_y) <= row_threshold:
                row.append(item)
                placed = True
                break

        if not placed:
            rows.append([item])

    rows.sort(key=lambda row: sum(it["yc"] for it in row) / len(row))

    result: list[list[str]] = []

    for row in rows:
        row.sort(key=lambda it: it["x0"])

        row_h = sum(it["h"] for it in row) / len(row)
        gaps = []

        for i in range(1, len(row)):
            gaps.append(row[i]["x0"] - row[i - 1]["x1"])

        positive_gaps = [g for g in gaps if g > 0]
        avg_gap = (
            sum(positive_gaps) / len(positive_gaps)
            if positive_gaps
            else row_h
        )

        # Larger than normal word gap becomes new cell.
        cell_gap_threshold = max(row_h * 1.15, avg_gap * 1.8, 8.0)

        cells = []
        current_parts = []
        prev_x1 = None

        for item in row:
            if prev_x1 is not None:
                gap = item["x0"] - prev_x1

                if gap > cell_gap_threshold:
                    cell = " ".join(current_parts).strip()
                    if cell:
                        cells.append(cell)
                    current_parts = []

            current_parts.append(item["text"])
            prev_x1 = item["x1"]

        if current_parts:
            cells.append(" ".join(current_parts).strip())

        cleaned = []

        for cell in cells:
            cell = _normalize_line_text(cell)
            cell = _aggressive_korean_spacing(
                _final_korean_spacing_cleanup(
                    _apply_korean_spacing_heuristics(cell)
                )
            )
            cell = _remove_leaked_brand_tokens(cell)
            cell = _fix_broken_pdf_symbols(cell)

            if cell:
                cleaned.append(cell)

        if cleaned:
            result.append(cleaned)

    return result


def _excel_quality_export_pdf_to_excel(
    self,
    source_pdf_path: str,
    output_xlsx_path: str,
    options: OfficeExportOptions | None = None,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is not installed. Run: pip install openpyxl"
        ) from exc

    options = options or OfficeExportOptions()
    excel_mode = getattr(options, "excel_mode", "lines")

    source_path = Path(source_pdf_path)
    output_path = Path(output_xlsx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open(str(source_path))

    try:
        page_indexes = options.page_range or list(range(len(pdf)))

        if not page_indexes:
            raise RuntimeError("No pages selected for Excel export.")

        wb = Workbook()

        # Remove default sheet; we will create sheets explicitly.
        default_ws = wb.active
        wb.remove(default_ws)

        if excel_mode == "table":
            for page_index in page_indexes:
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                ws = wb.create_sheet(title=f"Page {page_index + 1}")

                rows = _improved_group_words_to_table_rows_v2(page)

                ws.cell(row=1, column=1, value=f"Page {page_index + 1}")
                ws.cell(row=1, column=2, value="Coordinate Table Estimate")

                current_row = 2

                for row in rows:
                    for col_index, value in enumerate(row, start=1):
                        ws.cell(row=current_row, column=col_index, value=value)
                    current_row += 1

                _style_excel_sheet(ws)
                _autosize_excel_columns(ws)

        else:
            ws = wb.create_sheet(title="PDF Text")
            ws.append(["Page", "Line No", "Text"])

            row_no = 2

            for page_index in page_indexes:
                if page_index < 0 or page_index >= len(pdf):
                    continue

                page = pdf.load_page(page_index)
                lines = page.get_text("text").splitlines()

                line_no = 1

                for line in lines:
                    line = _fix_broken_pdf_symbols(line.strip())

                    if not line:
                        continue

                    line = _aggressive_korean_spacing(
                        _final_korean_spacing_cleanup(
                            _apply_korean_spacing_heuristics(line)
                        )
                    )
                    line = _remove_leaked_brand_tokens(line)
                    line = _fix_broken_pdf_symbols(line)

                    ws.cell(row=row_no, column=1, value=page_index + 1)
                    ws.cell(row=row_no, column=2, value=line_no)
                    ws.cell(row=row_no, column=3, value=line)

                    row_no += 1
                    line_no += 1

            _style_excel_sheet(ws)
            _autosize_excel_columns(ws, max_width=100)

        wb.save(str(output_path))

    finally:
        pdf.close()


try:
    OfficeExporter.export_pdf_to_excel = _excel_quality_export_pdf_to_excel
except NameError:
    pass

